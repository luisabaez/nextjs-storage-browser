"""
sampling.py - Interim record sampling for the ERP conversion validation effort.

Relationship config is taken directly from "Sampling Tables Relationship.xlsx".
For a chosen target table we select N random records, gather the child records
linked to those records through the business-key field defined in the sheet,
and write the whole sample set to an Excel workbook in the Sampling/ folder of
the data bucket.

Interim limitations (pending methodology confirmation with Ethree / Wanda):
  * Selection is random (ORDER BY NEWID()); runs are NOT reproducible yet.
  * Only DIRECT children of the target are gathered (no multi-level recursion).
  * Sample size is a manual input.
  * Table names are the exact MOCK14 names from the sheet; not yet parameterized
    by mock iteration.
"""

import io
import json
import math
import re
import unicodedata
import uuid
from collections import defaultdict
from datetime import datetime, date
from decimal import Decimal

import pyodbc
import openpyxl

SAMPLING_FOLDER = "Sampling/"
MANIFEST_FOLDER = "Sampling/_manifests/"
MAX_SAMPLE_SIZE = 5000          # guardrail
IN_CHUNK = 1000                 # values per IN(...) batch (SQL Server 2100 param cap)
URL_TTL_SECONDS = 3600

# The consolidated MOCK*_VW_TBL tables live in the Hacienda_ERP database, NOT
# the Hacienda_ERP_Test database the Lambda's connection string points at. The
# Lambda's SQL login has read access to it, so we reach the tables with
# database-qualified names ([Hacienda_ERP].[schema].[table]) rather than
# changing the global connection (which other actions need for _Test). A
# request may override this via source_db.
SOURCE_DATABASE = "Hacienda_ERP"

# ── Config generated from "Sampling Tables Relationship.xlsx" ──────────────────
TARGET_TABLES = [
    'FIN_911_BUDGET_BALANCE_MOCK14_VW_TBL',
    'FIN_AP_INVOICES_MOCK14_VW_TBL',
    'FIN_AP_INVOICES_SIFDE_MOCK14_VW_TBL',
    'FIN_AR_INVOICES_LINES_MOCK14_VW_TBL',
    'FIN_ASSETS_MOCK14_VW_TBL',
    'FIN_AWARDS_MOCK14_VW_TBL',
    'FIN_BUDGET_BALANCE_MOCK14_VW_TBL',
    'FIN_CUSTOMER_MOCK14_VW_TBL',
    'FIN_GL_BALANCE_911_MOCK14_VW_TBL',
    'FIN_GL_BALANCE_OPEN_PERIODS_911_MOCK14_VW_TBL',
    'FIN_GL_BALANCE_OPEN_PERIODS_MOCK14_VW_TBL',
    'FIN_GL_BALANCES_MOCK14_VW_TBL',
    'FIN_REVENUE_BUDGET_MOCK14_VW_TBL',
    'SCM_BPA_MOCK14_VW_TBL',
    'SCM_CONTRACTS_ASG_MOCK14_VW_TBL',
    'SCM_ITEMS_ASG_MOCK14_VW_TBL',
    'SCM_ITEMS_MOCK14_VW_TBL',
    'SCM_LOCATION_MOCK14_VW_CONVERTED',
    'SCM_PURCHASE_ORDERS_FINAL_MOCK14_VW_TBL',
    'SCM_PURCHASE_ORDERS_FINAL_911_MOCK14_VW_TBL',
    'SCM_PURCHASE_ORDERS_FINAL_RETIRO_MOCK14_VW_TBL',
    'SCM_REQ_HDR_MOCK14_VW_TBL',
    'SCM_SUPPLIER_MOCK14_VW_TBL',
    'SCM_SUPPLIER_ASG_MOCK13_VW_TBL',
    'HCM_PERSON_MOCK14_BU_VW_CONVERTED_TBL',
]

# (child_table, parent_table, link_field) - link_field is the business-key name;
# it is resolved to a real column at run time (see resolve_link_column).
RELATIONSHIPS = [
    ('FIN_AP_INV_LINES_MOCK14_VW_TBL', 'FIN_AP_INVOICES_MOCK14_VW_TBL', 'Invoice Number'),
    ('FIN_AP_INV_LINES_SIFDE_MOCK14_VW_TBL', 'FIN_AP_INVOICES_SIFDE_MOCK14_VW_TBL', 'Invoice Number'),
    ('FIN_AR_INVOICES_DISTRIBUTION_MOCK14_VW_tbl', 'FIN_AR_INVOICES_LINES_MOCK14_VW_TBL', 'Line Transactions Flexfield Segment 1 - Legacy Invoice Number'),
    ('FIN_ASSETS_DISTRIBUTION_MOCK14_VW_TBL', 'FIN_ASSETS_MOCK14_VW_TBL', 'Tag Number'),
    ('FIN_AWARD_BUDGET_MOCK14_VW_TBL', 'FIN_AWARDS_MOCK14_VW_TBL', 'Award Number'),
    ('FIN_AWARD_BUDGET_PERIOD_MOCK14_VW_TBL', 'FIN_AWARDS_MOCK14_VW_TBL', 'Award Number'),
    ('FIN_AWARD_CFDA_MOCK14_VW_TBL', 'FIN_AWARDS_MOCK14_VW_TBL', 'Award Number'),
    ('FIN_AWARD_FUNDING_ALLOCATION_MOCK14_VW_TBL', 'FIN_AWARDS_MOCK14_VW_TBL', 'Award Number'),
    ('FIN_AWARD_FUNDING_MOCK14_VW_TBL', 'FIN_AWARDS_MOCK14_VW_TBL', 'Award Number'),
    ('FIN_AWARD_FUNDING_SOURCE_MOCK14_VW_TBL', 'FIN_AWARDS_MOCK14_VW_TBL', 'Award Number'),
    ('FIN_AWARD_PROJECTS_MOCK14_VW_TBL', 'FIN_AWARDS_MOCK14_VW_TBL', 'Award Number'),
    ('FIN_CUSTOMER_CONTACT_MOCK14_VW_TBL', 'FIN_CUSTOMER_MOCK14_VW_TBL', 'Organization Name'),
    ('FIN_PRIFAS_CUSTOMER_MOCK14_VW_TBL', 'FIN_AWARDS_MOCK14_VW_TBL', 'Primary Sponsor Name'),
    ('FIN_PROJECT_TASKS_MOCK14_VW_TBL', 'FIN_PROJECTS_MOCK14_VW_TBL', 'Project Name'),
    ('FIN_PROJECTS_CLASS_MOCK14_VW_TBL', 'FIN_PROJECTS_MOCK14_VW_TBL', 'Project Name'),
    ('FIN_PROJECTS_MOCK14_VW_TBL', 'FIN_AWARD_PROJECTS_MOCK14_VW_TBL', 'Project Name'),
    ('FIN_PROJECTS_TEAM_MEMBERS_MOCK14_VW_TBL', 'FIN_AWARDS_MOCK14_VW_TBL', 'Project Name'),
    ('SCM_BPA_LINES_ATTACHMENT_MOCK14_VW_TBL', 'SCM_BPA_MOCK14_VW_TBL', 'BPA Number'),
    ('SCM_BPA_LINES_MOCK14_VW_TBL', 'SCM_BPA_MOCK14_VW_TBL', 'BPA Number'),
    ('SCM_CONTRACT_ATTACHMENTS_ASG_MOCK14_VW_TBL', 'SCM_CONTRACTS_ASG_MOCK14_VW_TBL', 'Contract Number'),
    ('SCM_CONTRACT_LINES_ASG_MOCK14_VW_TBL', 'SCM_CONTRACTS_ASG_MOCK14_VW_TBL', 'Contract Number'),
    ('SCM_CONTRACTS_PARTYCONTACTS_ASG_MOCK14_VW_TBL', 'SCM_CONTRACTS_ASG_MOCK14_VW_TBL', 'Contract Number'),
    ('SCM_INV_MANUFACTURER_MOCK14_VW_TBL', 'SCM_ITEMS_MOCK14_VW_TBL', 'Item'),
    ('SCM_ITEM_LOTS_MOCK14_VW_TBL', 'SCM_ITEMS_MOCK14_VW_TBL', 'Item'),
    ('SCM_ITEM_SUBINVENTORY_MOCK14_VW_TBL', 'SCM_ITEMS_MOCK14_VW_TBL', 'Item'),
    ('SCM_ITEMS_CATEGORY_MOCK14_VW_TBL', 'SCM_ITEMS_MOCK14_VW_TBL', 'Item'),
    ('SCM_ITEMS_LOCATOR_MOCK14_VW_TBL', 'SCM_ITEMS_MOCK14_VW_TBL', 'Item'),
    ('SCM_ITEMS_OHQ_MOCK14_VW_TBL', 'SCM_ITEMS_MOCK14_VW_TBL', 'Item'),
    ('SCM_ITEMS_RELATIONSHIPS_MOCK14_VW_TBL', 'SCM_ITEMS_MOCK14_VW_TBL', 'Item'),
    ('SCM_PS_ITEMS_CATEGORY_MOCK14_VW_TBL', 'SCM_PS_ITEMS_MOCK14_VW_TBL', 'Item'),
    ('SCM_PS_ITEMS_MOCK14_VW_TBL', 'SCM_PURCHASE_ORDERS_LINES_MOCK14_VW_TBL', 'Item'),
    # Purchase Orders — reconciled to the actual conversion-plan tables: the PO
    # header view is *_FINAL_* and the line/location/distribution views are
    # *_BY_BU_*, each split by source (main/PRIFAS, 911, RETIRO). All children
    # link to their source's header on the Order number.
    ('SCM_PURCHASE_ORDERS_LINES_FINAL_MOCK14_BY_BU_VW', 'SCM_PURCHASE_ORDERS_FINAL_MOCK14_VW_TBL', 'Order'),
    ('SCM_PURCHASE_ORDERS_LINE_LOCATIONS_MOCK14_BY_BU_VW', 'SCM_PURCHASE_ORDERS_FINAL_MOCK14_VW_TBL', 'Order'),
    ('SCM_PURCHASE_ORDERS_LINES_DISTRIBUTION_MOCK14_BY_BU_VW', 'SCM_PURCHASE_ORDERS_FINAL_MOCK14_VW_TBL', 'Order'),
    ('SCM_PURCHASE_ORDERS_LINES_911_MOCK14_BY_BU_VW', 'SCM_PURCHASE_ORDERS_FINAL_911_MOCK14_VW_TBL', 'Order'),
    ('SCM_PURCHASE_ORDERS_LINE_LOCATIONS_911_MOCK14_BY_BU_VW', 'SCM_PURCHASE_ORDERS_FINAL_911_MOCK14_VW_TBL', 'Order'),
    ('SCM_PURCHASE_ORDERS_LINES_DISTRIBUTION_911_MOCK14_BY_BU_VW', 'SCM_PURCHASE_ORDERS_FINAL_911_MOCK14_VW_TBL', 'Order'),
    ('SCM_PURCHASE_ORDERS_LINES_ATTACHMENTS_RET911_MOCK14_VW_TBL', 'SCM_PURCHASE_ORDERS_FINAL_911_MOCK14_VW_TBL', 'Order'),
    ('SCM_PURCHASE_ORDERS_LINES_RETIRO_MOCK14_BY_BU_VW', 'SCM_PURCHASE_ORDERS_FINAL_RETIRO_MOCK14_VW_TBL', 'Order'),
    ('SCM_PURCHASE_ORDERS_LINE_LOCATIONS_RETIRO_MOCK14_BY_BU_VW', 'SCM_PURCHASE_ORDERS_FINAL_RETIRO_MOCK14_VW_TBL', 'Order'),
    ('SCM_PURCHASE_ORDERS_LINES_DISTRIBUTION_RETIRO_MOCK14_BY_BU_VW', 'SCM_PURCHASE_ORDERS_FINAL_RETIRO_MOCK14_VW_TBL', 'Order'),
    ('SCM_REQ_DISTRIBUTION_MOCK14_VW_TBL', 'SCM_REQ_HDR_MOCK14_VW_TBL', 'Requisition Number'),
    ('SCM_REQ_LINE_MOCK14_VW_TBL', 'SCM_REQ_HDR_MOCK14_VW_TBL', 'Requisition Number'),
    ('SCM_SUPPLIER_ADDRESSES_ASG_MOCK14_VW_TBL', 'SCM_SUPPLIER_ASG_MOCK13_VW_TBL', 'Supplier Name'),
    ('SCM_SUPPLIER_ADDRESSES_MOCK14_VW_TBL', 'SCM_SUPPLIER_MOCK14_VW_TBL', 'Supplier Name'),
    ('SCM_SUPPLIER_BANK_ACCOUNTS_MOCK14_VW_TBL', 'SCM_SUPPLIER_MOCK14_VW_TBL', 'Supplier Name'),
    ('SCM_SUPPLIER_CLASSIFICATION_MOCK14_VW_TBL', 'SCM_SUPPLIER_MOCK14_VW_TBL', 'Supplier Name'),
    ('SCM_SUPPLIER_CONTACTS_MOCK14_VW_TBL', 'SCM_SUPPLIER_MOCK14_VW_TBL', 'Supplier Name'),
    ('SCM_SUPPLIER_SITE_ASG_MOCK14_VW_TBL', 'SCM_SUPPLIER_ASG_MOCK13_VW_TBL', 'Supplier Name'),
    ('SCM_SUPPLIER_SITE_ASSIG_ASG_MOCK14_VW_TBL', 'SCM_SUPPLIER_ASG_MOCK13_VW_TBL', 'Supplier Name'),
    ('SCM_SUPPLIER_SITE_ASSIG_MOCK14_VW_TBL', 'SCM_SUPPLIER_MOCK14_VW_TBL', 'Supplier Name'),
    ('SCM_SUPPLIER_SITE_MOCK14_VW_TBL', 'SCM_SUPPLIER_MOCK14_VW_TBL', 'Supplier Name'),
    ('SCM_SUPPLIER_CONTACTS_ASG_MOCK13_VW_TBL', 'SCM_SUPPLIER_ASG_MOCK13_VW_TBL', 'Supplier Name'),
    # HCM Person (#9): one entity with 8 sub-entities, each its own conversion
    # table, all linking to the Person on PERSON_NUMBER. (Person External
    # Identifier's HCM_EXTERNAL_IDENTIFIER_MOCK14_ALL is a different table shape
    # and isn't wired yet.)
    ('HCM_PERSON_ADDRESS_MOCK14_BU_VW_CONVERTED_TBL', 'HCM_PERSON_MOCK14_BU_VW_CONVERTED_TBL', 'PERSON_NUMBER'),
    ('HCM_PERSON_EMAIL_MOCK14_BU_VW_CONVERTED_TBL', 'HCM_PERSON_MOCK14_BU_VW_CONVERTED_TBL', 'PERSON_NUMBER'),
    ('HCM_PERSON_NAME_MOCK14_BU_VW_CONVERTED_TBL', 'HCM_PERSON_MOCK14_BU_VW_CONVERTED_TBL', 'PERSON_NUMBER'),
    ('HCM_PERSON_NID_MOCK14_BU_VW_CONVERTED_TBL', 'HCM_PERSON_MOCK14_BU_VW_CONVERTED_TBL', 'PERSON_NUMBER'),
    ('HCM_PERSON_ASSIGNMENT_MOCK14_BU_VW_CONVERTED_TBL', 'HCM_PERSON_MOCK14_BU_VW_CONVERTED_TBL', 'PERSON_NUMBER'),
    ('HCM_PERSON_SUPERVISOR_MOCK14_BU_VW_CONVERTED_TBL', 'HCM_PERSON_MOCK14_BU_VW_CONVERTED_TBL', 'PERSON_NUMBER'),
    ('HCM_EXTERNAL_BANK_ACCOUNT_MOCK14_BU_VW_CONVERTED_TBL', 'HCM_PERSON_MOCK14_BU_VW_CONVERTED_TBL', 'PERSON_NUMBER'),
]

_TARGET_SET = set(TARGET_TABLES)


# ── Config helpers ─────────────────────────────────────────────────────────────

def _short_name(table):
    """Friendly label: strip the _MOCKnn_VW[_TBL] suffix."""
    return re.sub(r'_MOCK\d+_VW(_TBL)?$', '', table, flags=re.IGNORECASE)


def _children_of(target):
    """[(child_table, link_field), ...] for the direct children of a target."""
    return [(c, link) for (c, p, link) in RELATIONSHIPS if p == target]


def list_targets():
    """Target tables + their direct children, for the picker UI."""
    out = []
    for t in TARGET_TABLES:
        children = [{"table": c, "display": _short_name(c), "link_field": link}
                    for (c, link) in _children_of(t)]
        out.append({
            "table": t,
            "display": _short_name(t),
            "child_count": len(children),
            "children": children,
        })
    out.sort(key=lambda r: r["display"])
    return {"ok": True, "count": len(out), "targets": out}


def list_relationships():
    """Full parent/child edge list + target tables, for client-side multi-level
    traversal (e.g. Awards -> Award Projects bridge -> Project Tasks). Config only,
    no DB/row data."""
    edges = [{"child": c, "parent": p, "link_field": link,
              "child_display": _short_name(c), "parent_display": _short_name(p)}
             for (c, p, link) in RELATIONSHIPS]
    return {"ok": True, "target_tables": list(TARGET_TABLES),
            "count": len(edges), "relationships": edges}


# ── Column resolution ──────────────────────────────────────────────────────────

def _normalize(s):
    return re.sub(r'[^a-z0-9]', '', (s or '').lower())


def _object_meta(cursor, db, table):
    """
    (schema, [columns]) for a table/view in `db`, via that database's
    INFORMATION_SCHEMA. Returns (None, []) when the object is not present.
    """
    cursor.execute(
        f"SELECT TABLE_SCHEMA, COLUMN_NAME FROM [{db}].INFORMATION_SCHEMA.COLUMNS "
        "WHERE TABLE_NAME = ? ORDER BY ORDINAL_POSITION",
        (table,),
    )
    rows = cursor.fetchall()
    if not rows:
        return None, []
    return rows[0][0], [r[1] for r in rows]


def _qualified(db, schema, table):
    return f"[{db}].[{schema}].[{table}]"


def resolve_link_column(columns, link_field):
    """
    Map a business-key name (e.g. "Invoice Number") to a real column name.

    Exact normalized match wins; otherwise the best substring match. Returns
    None when nothing matches confidently so the caller can flag it instead of
    guessing wrong.
    """
    norm_link = _normalize(link_field)
    if not norm_link:
        return None
    norm_cols = [(c, _normalize(c)) for c in columns]
    for c, nc in norm_cols:
        if nc == norm_link:
            return c
    # substring either direction, prefer the shortest column name (most specific)
    candidates = [c for c, nc in norm_cols if nc and (norm_link in nc or nc in norm_link)]
    if candidates:
        return min(candidates, key=len)
    return None


# ── Row / Excel helpers ────────────────────────────────────────────────────────

# openpyxl raises IllegalCharacterError on XML-illegal control chars; some converted
# item descriptions carry them (e.g. a stray \x1f after a description), so strip them
# from every string cell before writing.
_ILLEGAL_XLSX_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def _coerce(v):
    """Make a DB value safe for openpyxl cells (strips control chars openpyxl rejects)."""
    if isinstance(v, str):
        return _ILLEGAL_XLSX_RE.sub('', v)
    if v is None or isinstance(v, (int, float, bool, datetime, date)):
        return v
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (bytes, bytearray)):
        return v.hex()
    return _ILLEGAL_XLSX_RE.sub('', str(v))


def _fetch(cursor, sql, params=None):
    cursor.execute(sql, params or [])
    cols = [d[0] for d in cursor.description]
    rows = [list(r) for r in cursor.fetchall()]
    return cols, rows


_SHEET_BAD = re.compile(r'[\\/?*\[\]:]')


def _sheet_name(base, used):
    name = _SHEET_BAD.sub('_', base)[:31] or "Sheet"
    candidate, i = name, 1
    while candidate.lower() in used:
        suffix = f"_{i}"
        candidate = name[:31 - len(suffix)] + suffix
        i += 1
    used.add(candidate.lower())
    return candidate


def _add_sheet(wb, title, cols, rows):
    ws = wb.create_sheet(title=title)
    ws.append([str(c) for c in cols])
    for r in rows:
        ws.append([_coerce(v) for v in r])
    return ws


def _combine(parent_cols, parent_rows, p_link, child_cols, child_rows, c_link):
    """
    Flatten a parent/child relationship into one wide table joined on the link
    (unique-identifier) column. Child columns whose name already exists on the
    parent are dropped — repeating data is ignored so the shared key/fields show
    once — and exact-duplicate rows are collapsed. Parent rows with no matching
    child appear once with the child columns blank (left join), so the sampled
    records are never lost.

    Returns (combined_columns, combined_rows).
    """
    def _idx(cols, name):
        nl = name.lower()
        for i, c in enumerate(cols):
            if c.lower() == nl:
                return i
        return None

    p_idx = _idx(parent_cols, p_link)
    c_idx = _idx(child_cols, c_link)
    if p_idx is None or c_idx is None:
        # Links were already resolved; fail safe to parent-only if not found.
        return list(parent_cols), [list(r) for r in parent_rows]

    pset = {c.lower() for c in parent_cols}
    extra = [(i, c) for i, c in enumerate(child_cols) if c.lower() not in pset]
    combined_cols = list(parent_cols) + [c for _, c in extra]

    pmap = defaultdict(list)
    for pr in parent_rows:
        pmap[pr[p_idx]].append(pr)

    rows, matched = [], set()
    for cr in child_rows:
        parents = pmap.get(cr[c_idx])
        if not parents:
            continue
        matched.add(cr[c_idx])
        tail = [cr[i] for i, _ in extra]
        for pr in parents:
            rows.append(list(pr) + tail)

    blanks = [None] * len(extra)
    for pr in parent_rows:
        if pr[p_idx] not in matched:
            rows.append(list(pr) + blanks)

    seen, out = set(), []
    for r in rows:
        k = tuple('' if v is None else str(v) for v in r)
        if k in seen:
            continue
        seen.add(k)
        out.append(r)
    return combined_cols, out


def _combine_all(parent_cols, parent_rows, children_data):
    """
    Merge the parent and ALL its children into a single flattened sheet so the
    client can sort/filter every related record in one place. Each child record
    becomes one row carrying its parent's columns alongside; a leading "Source"
    column names the originating child table. Columns shared with the parent (and
    each child's link column) appear once — duplicated data is ignored — and
    child columns are unioned by name. Parents with no children at all are kept
    as a "(parent only)" row, and exact-duplicate rows are collapsed.

    children_data: list of (display, p_link, child_cols, child_rows, c_link).
    Returns (master_columns, master_rows).
    """
    def _idx(cols, name):
        nl = (name or "").lower()
        for i, c in enumerate(cols):
            if c.lower() == nl:
                return i
        return None

    pset = {c.lower() for c in parent_cols}
    master_cols = ["Source"] + list(parent_cols)
    seen_cols = {c.lower() for c in master_cols}
    for disp, p_link, ch_cols, ch_rows, c_link in children_data:
        cl_link = (c_link or "").lower()
        for c in ch_cols:
            cl = c.lower()
            if cl in pset or cl == cl_link or cl in seen_cols:
                continue
            master_cols.append(c)
            seen_cols.add(cl)

    pos = {c.lower(): i for i, c in enumerate(master_cols)}
    width = len(master_cols)
    n_parent = len(parent_cols)

    pmaps = {}
    def _pmap(p_idx):
        if p_idx not in pmaps:
            m = defaultdict(list)
            for ri, pr in enumerate(parent_rows):
                m[pr[p_idx]].append(ri)
            pmaps[p_idx] = m
        return pmaps[p_idx]

    rows, matched = [], set()
    for disp, p_link, ch_cols, ch_rows, c_link in children_data:
        p_idx = _idx(parent_cols, p_link)
        c_idx = _idx(ch_cols, c_link)
        if p_idx is None or c_idx is None:
            continue
        m = _pmap(p_idx)
        for cr in ch_rows:
            pris = m.get(cr[c_idx])
            if not pris:
                continue
            for ri in pris:
                matched.add(ri)
                pr = parent_rows[ri]
                row = [None] * width
                row[0] = disp
                for j in range(n_parent):
                    row[1 + j] = pr[j]
                for k, c in enumerate(ch_cols):
                    p = pos.get(c.lower())
                    if p is not None and p > n_parent:  # child-union cols only
                        row[p] = cr[k]
                rows.append(row)

    for ri, pr in enumerate(parent_rows):
        if ri not in matched:
            row = [None] * width
            row[0] = "(parent only)"
            for j in range(n_parent):
                row[1 + j] = pr[j]
            rows.append(row)

    seen_rows, out = set(), []
    for r in rows:
        k = tuple('' if v is None else str(v) for v in r)
        if k in seen_rows:
            continue
        seen_rows.add(k)
        out.append(r)
    return master_cols, out


# ── Main entry point ───────────────────────────────────────────────────────────

def run_sample(conn_str, s3, bucket, target_table, sample_size, actor="",
               source_db=None, combine_all=False):
    """
    Select `sample_size` random records from `target_table`, gather direct child
    records linked via each relationship's business key, write an Excel workbook
    to Sampling/ and return a summary + presigned download URL.

    Tables are read from `source_db` (default SOURCE_DATABASE) using
    database-qualified names, so the Lambda's _Test connection can reach the
    consolidated tables that live in Hacienda_ERP.
    """
    if target_table not in _TARGET_SET:
        return {"ok": False, "error": f"'{target_table}' is not a configured sampling target table."}
    try:
        sample_size = int(sample_size)
    except (TypeError, ValueError):
        return {"ok": False, "error": "sample_size must be a number."}
    if sample_size < 1:
        return {"ok": False, "error": "sample_size must be at least 1."}
    sample_size = min(sample_size, MAX_SAMPLE_SIZE)
    db = source_db or SOURCE_DATABASE

    with pyodbc.connect(conn_str) as conn:
        cur = conn.cursor()

        t_schema, target_cols = _object_meta(cur, db, target_table)
        if not target_cols:
            return {"ok": False, "error": f"Target table {target_table} not found in database {db}."}
        target_fq = _qualified(db, t_schema, target_table)

        # Total population (for context in the summary)
        cur.execute(f"SELECT COUNT(*) FROM {target_fq}")
        population = cur.fetchone()[0]

        # Random N rows
        p_cols, p_rows = _fetch(
            cur, f"SELECT TOP ({sample_size}) * FROM {target_fq} ORDER BY NEWID()"
        )
        col_idx = {c: i for i, c in enumerate(p_cols)}

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # drop default sheet
        used_sheets = set()

        child_summaries = []
        unresolved = []
        children = _children_of(target_table)

        if not children:
            # Standalone target (no children) — just its own records, like the
            # example workbook's "Budget Balance MC2" sheet.
            _add_sheet(wb, _sheet_name(_short_name(target_table), used_sheets), p_cols, p_rows)
        else:
            # Resolve + fetch every child once, then render either one combined
            # sheet per child, or (combine_all) a single merged master sheet.
            resolved = []
            for child, link_field in children:
                c_schema, child_cols = _object_meta(cur, db, child)
                if not child_cols:
                    unresolved.append({"child": child, "link_field": link_field,
                                       "reason": f"child table not found in {db}"})
                    continue
                p_link = resolve_link_column(p_cols, link_field)
                c_link = resolve_link_column(child_cols, link_field)
                if not p_link or not c_link:
                    unresolved.append({
                        "child": child, "link_field": link_field,
                        "reason": "could not resolve link column on "
                                  + ("parent" if not p_link else "child"),
                    })
                    continue

                key_values = sorted({
                    row[col_idx[p_link]] for row in p_rows
                    if row[col_idx[p_link]] is not None
                }, key=lambda x: str(x))

                child_fq = _qualified(db, c_schema, child)
                child_rows, ch_cols = [], child_cols
                for i in range(0, len(key_values), IN_CHUNK):
                    batch = key_values[i:i + IN_CHUNK]
                    placeholders = ",".join("?" * len(batch))
                    ch_cols, rows = _fetch(
                        cur,
                        f"SELECT * FROM {child_fq} WHERE [{c_link}] IN ({placeholders})",
                        batch,
                    )
                    child_rows.extend(rows)

                resolved.append({
                    "child": child, "display": _short_name(child),
                    "link_field": link_field, "p_link": p_link, "c_link": c_link,
                    "ch_cols": ch_cols, "child_rows": child_rows,
                })
                child_summaries.append({
                    "child": child, "display": _short_name(child),
                    "link_field": link_field, "parent_column": p_link,
                    "child_column": c_link, "row_count": len(child_rows),
                })

            if combine_all and resolved:
                # One master sheet: parent + all children unioned on the link
                # field, so the client can sort/filter all data in one place.
                children_data = [(r["display"], r["p_link"], r["ch_cols"],
                                  r["child_rows"], r["c_link"]) for r in resolved]
                m_cols, m_rows = _combine_all(p_cols, p_rows, children_data)
                _add_sheet(wb, _sheet_name(_short_name(target_table), used_sheets), m_cols, m_rows)
            else:
                # One combined, flattened sheet per child.
                for r in resolved:
                    comb_cols, comb_rows = _combine(
                        p_cols, p_rows, r["p_link"], r["ch_cols"], r["child_rows"], r["c_link"]
                    )
                    _add_sheet(wb, _sheet_name(r["display"], used_sheets), comb_cols, comb_rows)

            # If every child was unresolved, still write the parent sample so the
            # workbook isn't empty and the selection isn't lost.
            if not wb.sheetnames:
                _add_sheet(wb, _sheet_name(_short_name(target_table), used_sheets), p_cols, p_rows)

    # ── Write workbook + manifest to S3 ──
    stamp = datetime.utcnow().strftime("%Y%m%d-%H%M%S")
    stem = f"{_short_name(target_table)}_{stamp}_{uuid.uuid4().hex[:6]}"
    xlsx_key = f"{SAMPLING_FOLDER}{stem}.xlsx"
    filename = f"{stem}.xlsx"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    s3.put_object(
        Bucket=bucket, Key=xlsx_key, Body=buf.getvalue(),
        ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    manifest = {
        "target_table": target_table,
        "target_display": _short_name(target_table),
        "source_db": db,
        "combine_all": bool(combine_all),
        "requested_sample_size": sample_size,
        "selected_count": len(p_rows),
        "population": population,
        "children": child_summaries,
        "unresolved_links": unresolved,
        "child_total_rows": sum(c["row_count"] for c in child_summaries),
        "actor": actor or "",
        "created_at": datetime.utcnow().isoformat() + "Z",
        "xlsx_key": xlsx_key,
        "filename": filename,
        "method": "random (ORDER BY NEWID)",
    }
    s3.put_object(
        Bucket=bucket, Key=f"{MANIFEST_FOLDER}{stem}.json",
        Body=json.dumps(manifest, default=str), ContentType="application/json",
    )

    download_url = s3.generate_presigned_url(
        "get_object",
        Params={"Bucket": bucket, "Key": xlsx_key,
                "ResponseContentDisposition": f'attachment; filename="{filename}"'},
        ExpiresIn=URL_TTL_SECONDS,
    )

    return {"ok": True, "download_url": download_url, **manifest}


def list_runs(s3, bucket, limit=50):
    """Recent sampling runs, newest first, from the manifest folder."""
    runs = []
    paginator = s3.get_paginator("list_objects_v2")
    for page in paginator.paginate(Bucket=bucket, Prefix=MANIFEST_FOLDER):
        for obj in page.get("Contents", []):
            if not obj["Key"].endswith(".json"):
                continue
            try:
                body = s3.get_object(Bucket=bucket, Key=obj["Key"])["Body"].read()
                m = json.loads(body)
                m["_modified"] = obj["LastModified"].isoformat()
                runs.append(m)
            except Exception:
                continue
    runs.sort(key=lambda r: r.get("created_at", ""), reverse=True)
    return {"ok": True, "count": len(runs), "runs": runs[:limit]}


# ── Entity-file readiness (from SETUP_CONVERSION_PLAN_{mock}) ──────────────────
# The conversion plan defines, per entity/source/BU, the conversion table the
# generation scripts export to a CV_ file. This powers the Entity Files tab:
# what's expected, whether the table is populated, and (client-side) whether the
# file has been uploaded.
ENTITY_PLAN_COLUMNS = [
    'Pillar', 'Module', 'Entity', 'SubEntity', 'SOURCE', 'BU',
    'CONVERSION_TABLE_BU', 'CONVERSION_TABLE_SourceField', 'CONVERSION_TABLE_BU_Field',
    'Conversion_Table_Sourcefield_ForSampling', 'FileImportStatus', 'SourceFileName',
    'On Conversion Plan', 'ExcludedFromMock', 'RequiredFSCM',
]


def list_entity_plan(conn_str, mock='MOCK14', source_db=None, with_counts=False,
                     entity=None, raw=False):
    db = source_db or SOURCE_DATABASE
    plan_table = f"SETUP_CONVERSION_PLAN_{mock}"
    with pyodbc.connect(conn_str) as conn:
        cur = conn.cursor()
        cur.execute(
            f"SELECT COUNT(*) FROM [{db}].INFORMATION_SCHEMA.TABLES WHERE TABLE_NAME = ?",
            (plan_table,),
        )
        if cur.fetchone()[0] == 0:
            return {"ok": False, "error": f"{plan_table} not found in {db}"}

        # Which of the requested columns actually exist (schema varies by mock).
        cur.execute(
            f"SELECT COLUMN_NAME FROM [{db}].INFORMATION_SCHEMA.COLUMNS WHERE TABLE_NAME = ?",
            (plan_table,),
        )
        have = {r[0].lower() for r in cur.fetchall()}
        cols = [c for c in ENTITY_PLAN_COLUMNS if c.lower() in have]
        sel = ", ".join(f"[{c}]" for c in cols)
        # `raw` bypasses the has-a-converted-table filter so the FULL plan is visible
        # (diagnostic). Otherwise keep the generation-relevant filter. `entity` scopes
        # to one Entity. Qualifier columns (ExcludedFromMock/RequiredFSCM) are returned
        # when present so the caller can apply the client's in-scope rule.
        clauses, params = [], []
        if not raw:
            clauses.append("ISNULL([CONVERSION_TABLE_BU],'') <> ''")
            if 'enrichment_system' in have:
                clauses.append("ISNULL([ENRICHMENT_SYSTEM],'') <> 'Y'")
        if entity:
            clauses.append("[Entity] = ?")
            params.append(entity)
        where = ("WHERE " + " AND ".join(clauses)) if clauses else ""
        cur.execute(
            f"SELECT {sel} FROM [{db}].[dbo].[{plan_table}] {where} "
            f"ORDER BY [Pillar],[Module],[Entity],[SubEntity]", params
        )
        rows = cur.fetchall()
        out = []
        for r in rows:
            out.append({cols[i]: (str(r[i]).strip() if r[i] is not None else '') for i in range(len(cols))})

        counts = {}
        if with_counts:
            tables = sorted({rec.get('CONVERSION_TABLE_BU', '') for rec in out if rec.get('CONVERSION_TABLE_BU')})
            for t in tables:
                cur.execute(
                    f"SELECT COUNT(*) FROM [{db}].INFORMATION_SCHEMA.COLUMNS WHERE TABLE_NAME = ?",
                    (t,),
                )
                if cur.fetchone()[0] == 0:
                    counts[t] = None  # object missing
                    continue
                try:
                    cur.execute(f"SELECT COUNT(*) FROM [{db}].[dbo].[{t}]")
                    counts[t] = cur.fetchone()[0]
                except Exception:
                    counts[t] = None
            for rec in out:
                rec['tableRows'] = counts.get(rec.get('CONVERSION_TABLE_BU', ''))

    return {"ok": True, "mock": mock, "count": len(out), "rows": out,
            "planColumns": sorted(have)}


# ── Entity-file generation (server-side equivalent of the ConvertedFilesBySource
# scripts) ────────────────────────────────────────────────────────────────────
# For one entity, read its conversion tables from SETUP_CONVERSION_PLAN_{mock},
# split each by its source field (and BU field) and write one CV_ file per
# source[/BU] into Sampling/Generated/{mock}/{entity}/ — the raw parent+child
# set that feeds the merge + sampling. dry_run previews (counts, no writes).
GENERATED_FOLDER = "Sampling/Generated/"
GENERATED_MANIFEST_FOLDER = GENERATED_FOLDER + "_manifests/"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GEN_MAX_FILES = 400  # safety cap per run


def _safe_name(s):
    return re.sub(r'[^A-Za-z0-9]+', '_', str(s or '')).strip('_') or 'NA'


def _table_prefix(t):
    for s in ['_VW_CONVERTED_TBL', 'CONVERTED_VW', 'VW_CONVERTED', 'VW_TBL', 'VW_tbl', 'VW', 'TBL']:
        t = t.replace(s, '')
    return t.strip('_').rstrip('_')


def _clean_name(s):
    # Mirror the conversion scripts' CLEAN_STRING_V4 on plan field names: strip
    # accents, collapse internal whitespace, lower-case. So a plan field with
    # padding/accents still resolves to its column and its table is not skipped.
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode('ascii')
    return re.sub(r'\s+', ' ', s).strip().lower()


def _resolve_col(cols, name):
    nl = _clean_name(name)
    if not nl:
        return None
    for c in cols:
        if _clean_name(c) == nl:
            return c
    # last resort: alphanumeric-only match (punctuation differences)
    an = re.sub(r'[^a-z0-9]', '', nl)
    if an:
        for c in cols:
            if re.sub(r'[^a-z0-9]', '', _clean_name(c)) == an:
                return c
    return None


def _rows_to_xlsx_bytes(headers, rows):
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append([str(h) for h in headers])
    for r in rows:
        ws.append([_coerce(v) for v in r])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_gen_manifest(s3, bucket, mock, entity, subentity, db, actor,
                        out_prefix, generated, missing, capped=None, bu_filter=None,
                        empties=None):
    """Audit trail for one generation run: which CV_ files were written, from
    which conversion table / source / BU, their row counts, and who ran it and
    when. Lets a suspect generated file be traced back to its source later.
    Written alongside the files under Sampling/Generated/_manifests/."""
    stamp = datetime.utcnow().strftime("%Y%m%d-%H%M%S")
    key = f"{GENERATED_MANIFEST_FOLDER}{mock}/{_safe_name(entity)}/{stamp}_{uuid.uuid4().hex[:6]}.json"
    manifest = {
        "mock": mock,
        "entity": entity,
        "subentity": subentity or "",
        "source_db": db,
        "actor": actor or "",
        "created_at": datetime.utcnow().isoformat() + "Z",
        "folder": out_prefix,
        "file_count": len(generated),
        "total_rows": sum(g.get("rows", 0) for g in generated),
        "generated": generated,
        "missing": missing,
        "empties": empties or [],
    }
    if bu_filter:
        manifest["bu_filter"] = bu_filter
    if capped:
        manifest["capped"] = capped
    try:
        s3.put_object(Bucket=bucket, Key=key,
                      Body=json.dumps(manifest, default=str),
                      ContentType="application/json")
    except Exception:
        return None
    return key


def _empty_entry(table, prefix, bu_filter, reason):
    """One record for an expected conversion table that produced no file this run,
    so the UI can distinguish empty / not-built from never-generated (both would
    otherwise just look 'missing'). reason ∈ not_built | no_source_field |
    empty_table | no_rows_for_bu."""
    return {"table": table, "prefix": prefix, "source": "",
            "bu": bu_filter or "", "rows": 0, "reason": reason}


# Entities whose source/BU split value is a 7-digit ledger segment (e.g. 0150000)
# whose first 3 digits are the agency/BU. Their per-BU filter must match a
# 3-digit BU against that agency prefix, not the exact 7-digit value.
_LEDGER_SEGMENT_ENTITIES = {'gl balances', 'gl budget balances'}

# Entities whose conversion tables are gated by the plan's "On Conversion Plan"
# flag: only tables marked Y are generated, checked live each run. Scoped to the
# sampling-wired entities so an all-N entity elsewhere (e.g. BPA) is unaffected.
# Finance Location was added by request while its rows are still flagged N, so it
# is intentionally NOT gated — it generates regardless of the flag.
_ON_PLAN_ENTITIES = {'gl balances', 'gl budget balances'}


# Some entities key their conversion tables by a named source SYSTEM rather than a
# numeric BU (AR Customer is tagged SALUD / SIFDE, not 071 / 081). Map the name to
# its BU so a per-BU generation call still matches. Add new named sources here.
_SOURCE_BU_MAP = {'SALUD': '071', 'SIFDE': '081'}

# Confirmed-orphan source families. A legacy source system (e.g. 911) whose master
# rows were converted under a placeholder BU (000) but that the client has confirmed
# belongs to a real agency. We emit that source's master rows + its children (joined
# on the relationship key, since the children carry only the placeholder BU) re-tagged
# to the real agency, so a per-BU run for that agency surfaces them and the client's
# per-source split names the workbook "<Entity> BU <real_bu> <source>" (e.g.
# "Supplier BU 045 911"). Keyed on the master's SOURCE value, not its stored BU.
_ORPHAN_SOURCE_FAMILIES = [
    {'entity': 'Supplier', 'source': '911', 'real_bu': '045'},
]


def _bu_matches(bu_filter, source, bu, ledger_segment=False):
    """Whether a source/BU split belongs to the requested BU. Exact match for
    normal entities; a named source system maps to its BU (SALUD -> 071); a 3-digit BU
    also matches the agency prefix of a 7-digit segment value (0150000 -> 015,
    0450121 -> 045)."""
    if bu_filter in (source, bu):
        return True
    # A named source system (SALUD -> 071) resolves to its BU, but only for a source-
    # keyed split with no numeric BU of its own — so entities that use SALUD / SIFDE as
    # a source system while carrying a real BU aren't mis-matched to 071 / 081. Its
    # sub-sources resolve via the first delimited token (SALUD-FACTURASALCOBRO -> SALUD).
    if not bu and source:
        su = str(source).strip().upper()
        first = re.split(r'[^A-Z0-9]', su)[0] if su else ''
        if _SOURCE_BU_MAP.get(su) == bu_filter or (first and _SOURCE_BU_MAP.get(first) == bu_filter):
            return True
    # A 7-digit ledger-style segment carries the 3-digit BU as its prefix. The form is
    # self-identifying, so match it for ANY entity (e.g. the Purchase Order
    # DISTRIBUTION_BY_BU child, whose source is the 7-digit segment 0160000), not just
    # the ledger-segment entities.
    want = bu_filter.lstrip('0') or '0'
    for v in (source, bu):
        v = str(v).strip()
        if v.isdigit() and len(v) == 7 and (v[:3].lstrip('0') or '0') == want:
            return True
    # Inventory keys the BU inside an Organization code (INV_016651 / INV_010RCV -> 016 /
    # 010) or as the leading token of an OHQ "016 AGENCY NAME" BU field — extract the
    # embedded 3-digit BU. INV_ is unique to Inventory; the "NNN <text>" form (digits +
    # whitespace) doesn't collide with clean BUs, named sources, or 7-digit segments.
    for v in (source, bu):
        s = str(v).strip()
        m = re.match(r'INV_?(\d{3})', s, re.I) or re.match(r'(\d{3})\s', s)
        if m and (m.group(1).lstrip('0') or '0') == want:
            return True
    return False


def _linked_child_files(cur, db, table, prefix, bu_filter, ledger_segment, plans):
    """Generate a LINK-ONLY child (one with no agency column of its own, e.g.
    Requisition Line / Distribution) by joining it to its parent on the relationship
    key and inheriting the parent's source/BU. This lets the tables stay unchanged:
    the child follows whatever keying the parent uses (a legacy-system name today, or
    a numeric BU once the parent's BU_Field is set), and lands in the parent's merge
    group so the client links it by key like any other child.

    Returns None if `table` is not a resolvable link-only child (caller falls back to
    the normal 'no source field' handling); otherwise a list of
    {"file","source","bu","headers","rows"} — one per parent source/BU combo that
    passes bu_filter (headers/rows are ready to write)."""
    # Match the relationship by BASE table identity: RELATIONSHIPS names tables as
    # _VW_TBL while the plan (and the populated data) use the entity's conversion table
    # (_VW_CONVERTED), so an exact-name match misses link-only children like Item
    # Category. Resolve the parent's REAL table from the plan so the join hits real rows.
    parent = next(((p, lf) for (c, p, lf) in RELATIONSHIPS if _base_table(c) == _base_table(table)), None)
    if not parent:
        return None
    rel_parent, link_field = parent
    p_plan = next((pl for pl in plans if _base_table((pl[0] or '').strip()) == _base_table(rel_parent)), None)
    if not p_plan:
        return None
    p_table = (p_plan[0] or '').strip()
    _, c_cols = _object_meta(cur, db, table)
    c_link = resolve_link_column(c_cols, link_field) if c_cols else None
    p_schema, p_cols = _object_meta(cur, db, p_table)
    if not (c_cols and c_link and p_cols):
        return None
    p_src_field = (p_plan[1] or '').strip()
    p_src = _resolve_col(p_cols, p_src_field)
    p_link = resolve_link_column(p_cols, link_field)
    if not (p_src and p_link):
        return None
    p_bu_field = (p_plan[2] or '').strip()
    p_bu = _resolve_col(p_cols, p_bu_field) if p_bu_field else None
    if p_bu and p_bu.lower() == p_src.lower():
        p_bu = None
    p_fq = f"[{db}].[dbo].[{p_table}]"
    c_fq = f"[{db}].[dbo].[{table}]"
    if p_bu:
        cur.execute(f"SELECT DISTINCT [{p_src}],[{p_bu}] FROM {p_fq} WHERE [{p_src}] <> ?", (p_src_field,))
    else:
        cur.execute(f"SELECT DISTINCT [{p_src}] FROM {p_fq} WHERE [{p_src}] <> ?", (p_src_field,))
    combos = cur.fetchall()
    out = []
    for pc in combos:
        raw_source = pc[0]
        if raw_source is None:
            continue
        source = str(raw_source).strip()
        has_bu = bool(p_bu) and len(pc) > 1 and pc[1] is not None
        raw_bu = pc[1] if has_bu else None
        bu = str(raw_bu).strip() if has_bu else ''
        if bu_filter and not _bu_matches(bu_filter, source, bu, ledger_segment):
            continue
        if has_bu:
            cur.execute(f"SELECT DISTINCT [{p_link}] FROM {p_fq} WHERE [{p_src}]=? AND [{p_bu}]=?", (raw_source, raw_bu))
        else:
            cur.execute(f"SELECT DISTINCT [{p_link}] FROM {p_fq} WHERE [{p_src}]=?", (raw_source,))
        keys = [r[0] for r in cur.fetchall() if r[0] is not None]
        fname = f"CV_{prefix}__{_safe_name(source)}" + (f"_{_safe_name(bu)}" if has_bu else "") + ".xlsx"
        headers, rows = list(c_cols), []
        for i in range(0, len(keys), IN_CHUNK):
            batch = keys[i:i + IN_CHUNK]
            ph = ",".join("?" * len(batch))
            cur.execute(f"SELECT * FROM {c_fq} WHERE [{c_link}] IN ({ph})", batch)
            headers = [d[0] for d in cur.description]
            rows.extend(cur.fetchall())
        out.append({"file": fname, "source": source, "bu": bu, "headers": headers, "rows": rows})
    return out


def _base_table(t):
    """Strip the mock/view suffix so a plan table (…_MOCK14_CONVERTED_VW) and its
    RELATIONSHIPS name (…_MOCK14_VW_TBL) compare equal on their base identity."""
    return re.sub(r'_MOCK\d+.*$', '', (t or '').strip(), flags=re.I).upper()


def _orphan_source_family_files(cur, db, plans, family):
    """Emit a confirmed-orphan source family (see _ORPHAN_SOURCE_FAMILIES): the master
    rows for one legacy source value plus every direct child (joined on the relationship
    key), all re-tagged to the agency the client assigned the source to. Returns a list
    of {file, table, source, bu, headers, rows}, or [] when not applicable to `plans`."""
    src_val, real_bu = family['source'], family['real_bu']
    # Master = the plan table keyed by a SOURCE field distinct from its BU field
    # (children key both on 'BU'); it is the RELATIONSHIPS parent of the children.
    master = next(((t, sf) for (t, sf, bf) in plans
                   if t and sf and (bf or '').strip().lower() != (sf or '').strip().lower()), None)
    if not master:
        return []
    m_table, m_src_field = (master[0] or '').strip(), (master[1] or '').strip()
    _, m_cols = _object_meta(cur, db, m_table)
    m_src = _resolve_col(m_cols, m_src_field) if m_cols else None
    if not m_src:
        return []
    m_fq = f"[{db}].[dbo].[{m_table}]"
    cur.execute(f"SELECT * FROM {m_fq} WHERE [{m_src}] = ?", (src_val,))
    m_headers = [d[0] for d in cur.description]
    m_rows = cur.fetchall()
    if not m_rows:
        return []
    out = [{"file": f"CV_{_table_prefix(m_table)}__{_safe_name(src_val)}_{_safe_name(real_bu)}.xlsx",
            "table": m_table, "source": src_val, "bu": real_bu, "headers": m_headers, "rows": list(m_rows)}]
    # Children link to the master on the relationship key; collect those keys from the
    # master rows so each child is filtered to exactly this source's suppliers.
    mbase = _base_table(m_table)
    for (t, _sf, _bf) in plans:
        t = (t or '').strip()
        if not t or _base_table(t) == mbase:
            continue
        link = next((lf for (c, p, lf) in RELATIONSHIPS
                     if _base_table(c) == _base_table(t) and _base_table(p) == mbase), None)
        if not link:
            continue  # not a direct child of this master (e.g. an ASG-parented table)
        _, c_cols = _object_meta(cur, db, t)
        c_link = resolve_link_column(c_cols, link) if c_cols else None
        m_link = resolve_link_column(m_headers, link)
        if not (c_link and m_link):
            continue
        li = m_headers.index(m_link)
        keys = sorted({r[li] for r in m_rows if r[li] is not None}, key=str)
        c_fq = f"[{db}].[dbo].[{t}]"
        c_headers, c_rows = list(c_cols), []
        for i in range(0, len(keys), IN_CHUNK):
            batch = keys[i:i + IN_CHUNK]
            ph = ",".join("?" * len(batch))
            cur.execute(f"SELECT * FROM {c_fq} WHERE [{c_link}] IN ({ph})", batch)
            c_headers = [d[0] for d in cur.description]
            c_rows.extend(cur.fetchall())
        out.append({"file": f"CV_{_table_prefix(t)}__{_safe_name(src_val)}_{_safe_name(real_bu)}.xlsx",
                    "table": t, "source": src_val, "bu": real_bu, "headers": c_headers, "rows": c_rows})
    return out


# HCM Person sub-entities (label, table stem with {m}=mock). Master first; the rest
# link to it by PERSON_NUMBER. Mirrors HCM_PERSON_SUB on the client.
_HCM_PERSON_SHEETS = [
    ('Person', 'HCM_PERSON_{m}_BU_VW_CONVERTED_TBL'),
    ('Person Address', 'HCM_PERSON_ADDRESS_{m}_BU_VW_CONVERTED_TBL'),
    ('Person Email', 'HCM_PERSON_EMAIL_{m}_BU_VW_CONVERTED_TBL'),
    ('Person Name', 'HCM_PERSON_NAME_{m}_BU_VW_CONVERTED_TBL'),
    ('Person NID', 'HCM_PERSON_NID_{m}_BU_VW_CONVERTED_TBL'),
    ('Assignment', 'HCM_PERSON_ASSIGNMENT_{m}_BU_VW_CONVERTED_TBL'),
    ('Supervisor', 'HCM_PERSON_SUPERVISOR_{m}_BU_VW_CONVERTED_TBL'),
    ('External Bank Account', 'HCM_EXTERNAL_BANK_ACCOUNT_{m}_BU_VW_CONVERTED_TBL'),
]


def sample_hcm_person(conn_str, sources, sample_size, mock='MOCK14', source_db=None):
    """Server-side HCM Person sampling by source system (ATTRIBUTE1). Picks
    `sample_size` random people whose ATTRIBUTE1 is in `sources`, then pulls only
    those people's rows from each Person sub-entity (linked on PERSON_NUMBER), so the
    browser can build its report without loading the full population. `sources` is a
    comma list; several published together (KRONOSPOL + ADPPOLICIA) are passed at once."""
    db = source_db or SOURCE_DATABASE
    srcs = [s.strip() for s in (sources or '').split(',') if s.strip()]
    if not srcs:
        return {"ok": False, "error": "sources required"}
    try:
        n_req = int(sample_size)
    except (TypeError, ValueError):
        return {"ok": False, "error": "sample_size must be a number"}
    master_table = _HCM_PERSON_SHEETS[0][1].format(m=mock)
    with pyodbc.connect(conn_str) as conn:
        cur = conn.cursor()
        m_schema, m_cols = _object_meta(cur, db, master_table)
        if not m_cols:
            return {"ok": False, "error": f"{master_table} not found in {db}"}
        src_col = _resolve_col(m_cols, 'ATTRIBUTE1')
        pn_col = resolve_link_column(m_cols, 'PERSON_NUMBER')
        if not (src_col and pn_col):
            return {"ok": False, "error": "ATTRIBUTE1 or PERSON_NUMBER column not found on the Person table"}
        m_fq = _qualified(db, m_schema, master_table)
        ph = ",".join("?" * len(srcs))
        cur.execute(f"SELECT COUNT(*) FROM {m_fq} WHERE [{src_col}] IN ({ph})", srcs)
        population = cur.fetchone()[0]
        sheets, keys = [], []
        n = max(1, min(n_req, population)) if population else 0
        if n:
            cur.execute(f"SELECT TOP ({n}) * FROM {m_fq} WHERE [{src_col}] IN ({ph}) ORDER BY NEWID()", srcs)
            headers = [d[0] for d in cur.description]
            rows = [[_coerce(v) for v in r] for r in cur.fetchall()]
            pn_i = headers.index(pn_col)
            keys = sorted({r[pn_i] for r in rows if r[pn_i] is not None}, key=lambda x: str(x))
            sheets.append({"label": _HCM_PERSON_SHEETS[0][0], "table": master_table, "headers": headers, "rows": rows})
            for label, stem in _HCM_PERSON_SHEETS[1:]:
                table = stem.format(m=mock)
                c_schema, c_cols = _object_meta(cur, db, table)
                if not c_cols:
                    continue
                c_pn = resolve_link_column(c_cols, 'PERSON_NUMBER')
                if not c_pn:
                    continue
                c_fq = _qualified(db, c_schema, table)
                c_headers, c_rows = list(c_cols), []
                for i in range(0, len(keys), IN_CHUNK):
                    batch = keys[i:i + IN_CHUNK]
                    cph = ",".join("?" * len(batch))
                    cur.execute(f"SELECT * FROM {c_fq} WHERE [{c_pn}] IN ({cph})", batch)
                    c_headers = [d[0] for d in cur.description]
                    c_rows.extend([[_coerce(v) for v in r] for r in cur.fetchall()])
                sheets.append({"label": label, "table": table, "headers": c_headers, "rows": c_rows})
    return {"ok": True, "population": population, "sample_size": len(keys),
            "person_key": pn_col, "sources": srcs, "sheets": sheets}


def _derive_bu(source, bu):
    """The 3-digit BU a plan row belongs to, derived from its source/BU values — the
    inverse of _bu_matches (which tests a known BU). Handles a clean numeric BU, a 7-digit
    ledger segment (0500000 -> 050), an INV_<bu> org code, a "NNN name" OHQ value, and a
    named source system via _SOURCE_BU_MAP (SALUD -> 071). Returns None when unmappable
    (e.g. PRIFAS, which spans many agencies)."""
    vals = [v for v in (bu, source) if v is not None and str(v).strip() != '']
    for v in vals:
        s = str(v).strip()
        if s.isdigit() and 1 <= len(s) <= 3:
            return s.zfill(3)
    for v in (source, bu):
        s = str(v).strip() if v is not None else ''
        if s.isdigit() and len(s) == 7:
            return s[:3]
    for v in (source, bu):
        s = str(v).strip() if v is not None else ''
        m = re.match(r'INV_?(\d{3})', s, re.I) or re.match(r'(\d{3})\s', s)
        if m:
            return m.group(1)
    for v in (source, bu):
        s = str(v).strip().upper() if v is not None else ''
        if not s:
            continue
        first = re.split(r'[^A-Z0-9]', s)[0]
        if _SOURCE_BU_MAP.get(s):
            return _SOURCE_BU_MAP[s]
        if first and _SOURCE_BU_MAP.get(first):
            return _SOURCE_BU_MAP[first]
    return None


def entity_bu_coverage(conn_str, mock='MOCK14', source_db=None):
    """Which BUs actually have converted data for each master entity — the LIVE readiness
    source, so the app needn't rely on the prior-run agency spreadsheet. For every plan
    master (Included + RequiredFSCM=Y whose conversion table base-matches a sampling
    TARGET), read its converted table's distinct source/BU values and derive the 3-digit
    BUs present. Returned keyed by the table's base identity (SCM_ITEMS, ...), which the
    client resolves each entity tab to."""
    db = source_db or SOURCE_DATABASE
    plan_table = f"SETUP_CONVERSION_PLAN_{mock}"
    targets_base = {_base_table(t) for t in TARGET_TABLES}
    coverage = {}
    with pyodbc.connect(conn_str) as conn:
        cur = conn.cursor()
        cur.execute(
            f"SELECT DISTINCT [CONVERSION_TABLE_BU],[CONVERSION_TABLE_SourceField],"
            f"[CONVERSION_TABLE_BU_Field] FROM [{db}].[dbo].[{plan_table}] "
            f"WHERE ISNULL([ExcludedFromMock],'')='Included' AND ISNULL([RequiredFSCM],'')='Y' "
            f"AND ISNULL([CONVERSION_TABLE_BU],'')<>''")
        seen = set()
        for (table, srcf, buf) in cur.fetchall():
            table = (table or '').strip()
            base = _base_table(table)
            if base not in targets_base or table in seen:
                continue
            seen.add(table)
            schema, cols = _object_meta(cur, db, table)
            if not cols:
                continue
            src_col = _resolve_col(cols, (srcf or '').strip()) if (srcf or '').strip() else None
            bu_col = _resolve_col(cols, (buf or '').strip()) if (buf or '').strip() else None
            sel = [c for c in (src_col, bu_col) if c]
            if not sel:
                continue
            fq = _qualified(db, schema, table)
            si = sel.index(src_col) if src_col in sel else -1
            bi = sel.index(bu_col) if bu_col in sel else -1
            cur.execute(f"SELECT DISTINCT {','.join('['+c+']' for c in sel)} FROM {fq}")
            bset = coverage.setdefault(base, set())
            for row in cur.fetchall():
                b = _derive_bu(row[si] if si >= 0 else None, row[bi] if bi >= 0 else None)
                if b:
                    bset.add(b)
    return {"ok": True, "coverage": {k: sorted(v) for k, v in coverage.items()}}


def sample_entity_by_bu(conn_str, entity, bu, sample_size, mock='MOCK14', source_db=None,
                        z=None, e=None, p=None):
    """Server-side per-BU sampling for any entity whose population is too large to build
    in the browser (e.g. Purchase Orders BU 081). Reads the entity's conversion plan to
    find the root (target) table and how it keys the BU, samples `sample_size` root rows
    for that BU, then fetches only those rows' children (linked via the relationship
    graph). Returns the sampled sheets so the browser builds its report from a small set.

    If tier params (z,e,p) are supplied the sample is sized here via Cochran against the
    LIVE population, so the caller needn't know the population up front (one round trip)."""
    db = source_db or SOURCE_DATABASE
    try:
        n_req = int(sample_size)
    except (TypeError, ValueError):
        return {"ok": False, "error": "sample_size must be a number"}
    bu = str(bu).strip()
    ledger = (entity or '').strip().lower() in _LEDGER_SEGMENT_ENTITIES
    plan_table = f"SETUP_CONVERSION_PLAN_{mock}"
    with pyodbc.connect(conn_str) as conn:
        cur = conn.cursor()
        where = ("ISNULL([CONVERSION_TABLE_BU],'') <> '' AND ISNULL([ENRICHMENT_SYSTEM],'') <> 'Y' "
                 "AND ISNULL([CONVERSION_TABLE_SourceField],'') <> '' AND [Entity] = ?")
        cur.execute(f"SELECT DISTINCT [CONVERSION_TABLE_BU],[CONVERSION_TABLE_SourceField],"
                    f"[CONVERSION_TABLE_BU_Field] FROM [{db}].[dbo].[{plan_table}] WHERE {where}", (entity,))
        plans = cur.fetchall()
        # Find the root table (a sampling target) whose distinct source/BU combos match
        # the requested BU. Build the WHERE that selects that BU's root rows.
        root = None
        r_src_col = r_bu_col = None
        conds, params = [], []
        for (table, src_field, bu_field) in plans:
            table = (table or '').strip()
            if table not in _TARGET_SET:
                continue
            _, cols = _object_meta(cur, db, table)
            if not cols:
                continue
            src_field = (src_field or '').strip()
            src_col = _resolve_col(cols, src_field)
            if not src_col:
                continue
            bu_col = _resolve_col(cols, (bu_field or '').strip()) if (bu_field or '').strip() else None
            if bu_col and bu_col.lower() == src_col.lower():
                bu_col = None
            fq = f"[{db}].[dbo].[{table}]"
            if bu_col:
                cur.execute(f"SELECT DISTINCT [{src_col}],[{bu_col}] FROM {fq} WHERE [{src_col}] <> ?", (src_field,))
            else:
                cur.execute(f"SELECT DISTINCT [{src_col}] FROM {fq} WHERE [{src_col}] <> ?", (src_field,))
            for c in cur.fetchall():
                if c[0] is None:
                    continue
                rs = str(c[0]).strip()
                has_bu = bool(bu_col) and len(c) > 1 and c[1] is not None
                rb = str(c[1]).strip() if has_bu else ''
                if not _bu_matches(bu, rs, rb, ledger):
                    continue
                if has_bu:
                    conds.append(f"([{src_col}]=? AND [{bu_col}]=?)"); params += [c[0], c[1]]
                else:
                    conds.append(f"[{src_col}]=?"); params += [c[0]]
            if conds:
                root, r_src_col, r_bu_col = table, src_col, bu_col
                break
        if not root or not conds:
            return {"ok": False, "error": f"no root table matched BU {bu} for {entity}"}
        r_schema, _ = _object_meta(cur, db, root)
        r_fq = _qualified(db, r_schema, root)
        wc = " OR ".join(conds)
        cur.execute(f"SELECT COUNT(*) FROM {r_fq} WHERE {wc}", params)
        population = cur.fetchone()[0]
        if z is not None and e is not None and p is not None:
            n = _cochran_n(population, z, e, p)
        else:
            n = max(1, min(n_req, population)) if population else 0
        sheets = []
        if n:
            cur.execute(f"SELECT TOP ({n}) * FROM {r_fq} WHERE {wc} ORDER BY NEWID()", params)
            r_headers = [d[0] for d in cur.description]
            r_rows = [[_coerce(v) for v in row] for row in cur.fetchall()]
            sheets.append({"label": _short_name(root), "table": root, "headers": r_headers,
                           "rows": r_rows, "population": population})
            for child, link_field in _children_of(root):
                c_schema, c_cols = _object_meta(cur, db, child)
                if not c_cols:
                    continue
                p_link = resolve_link_column(r_headers, link_field)
                c_link = resolve_link_column(c_cols, link_field)
                if not (p_link and c_link):
                    continue
                pi = r_headers.index(p_link)
                keys = sorted({row[pi] for row in r_rows if row[pi] is not None}, key=lambda x: str(x))
                c_fq = _qualified(db, c_schema, child)
                # True population = every child row for ALL of this BU's root rows (not just
                # the sampled ones), for the Relationships/Sizing "Population rows" column.
                cur.execute(f"SELECT COUNT(*) FROM {c_fq} WHERE [{c_link}] IN "
                            f"(SELECT [{p_link}] FROM {r_fq} WHERE {wc})", params)
                c_pop = cur.fetchone()[0]
                c_headers, c_rows = list(c_cols), []
                for i in range(0, len(keys), IN_CHUNK):
                    batch = keys[i:i + IN_CHUNK]
                    ph = ",".join("?" * len(batch))
                    cur.execute(f"SELECT * FROM {c_fq} WHERE [{c_link}] IN ({ph})", batch)
                    c_headers = [d[0] for d in cur.description]
                    c_rows.extend([[_coerce(v) for v in row] for row in cur.fetchall()])
                sheets.append({"label": _short_name(child), "table": child, "headers": c_headers,
                               "rows": c_rows, "population": c_pop})
    # The root's unique-id column (what children link on) — for the report's highlight.
    root_key = None
    for child, link_field in _children_of(root):
        rk = resolve_link_column(sheets[0]["headers"], link_field) if sheets else None
        if rk:
            root_key = rk
            break
    return {"ok": True, "entity": entity, "bu": bu, "root": root, "population": population,
            "sample_size": (len(sheets[0]["rows"]) if sheets else 0),
            "root_key": root_key, "sheets": sheets}


# Inventory sample scope (Entity_Hierarchy.docx): Items master + Item Category + Item OHQ,
# all linked on 'Item'. The master keys the BU inside its Organization code (INV_016651),
# resolved by _bu_matches. Kept explicit (like _HCM_PERSON_SHEETS) so the sample is exactly
# these three tables — not every SCM_ITEMS graph child.
_INVENTORY_ROOT = 'SCM_ITEMS_{m}_VW_CONVERTED'
_INVENTORY_CHILDREN = [
    ('Item Category', 'SCM_ITEMS_CATEGORY_{m}_VW_CONVERTED'),
    ('Item OHQ', 'SCM_ITEMS_OHQ_{m}_VW_CONVERTED'),
]
_INVENTORY_ORG_FIELD = 'Organization'
_INVENTORY_LINK = 'Item'


def _cochran_n(N, z, e, p):
    """Cochran attribute sample size with finite-population correction (Framework V2),
    mirroring the client's computeSampleSize: n = N*Z^2*p*(1-p) / [e^2*(N-1) + Z^2*p*(1-p)],
    rounded up, floored at 1, capped at N. Used to size EACH inventory organization on its
    own population so every org gets at least one sampled item."""
    try:
        N = int(N); z = float(z); e = float(e); p = float(p)
    except (TypeError, ValueError):
        return 0
    if N <= 0:
        return 0
    zpq = z * z * p * (1 - p)
    n = (N * zpq) / (e * e * (N - 1) + zpq)
    return min(N, max(1, math.ceil(n)))


def sample_inventory_by_bu(conn_str, bu, sample_size, mock='MOCK14', source_db=None,
                           z=None, e=None, p=None, orgs=None):
    """Server-side Inventory sampling for one BU, kept in a SINGLE per-BU workbook.

    A BU can hold several inventory organizations (e.g. 024 = INV_024071/131/501/581/651).
    The client requires a sample drawn from EACH organization, so when tier params (z,e,p)
    are supplied every org is sized on its OWN population via Cochran (>=1 each) and the
    draws are unioned — no organization is left unrepresented. Without tier params it falls
    back to a single TOP(n) across all of the BU's orgs. Item Category + Item OHQ children
    are then pulled for the sampled items (linked on Item); the master population reported
    stays the total of the orgs sampled.

    `orgs` optionally restricts sampling to specific Organization codes (a list or a
    comma-separated string) — used to split a BU into separate workbooks by organization
    (e.g. 045 = INV_045321/652/655 in one file and INV_045656 (911) in its own)."""
    db = source_db or SOURCE_DATABASE
    try:
        n_req = int(sample_size)
    except (TypeError, ValueError):
        return {"ok": False, "error": "sample_size must be a number"}
    stratify = z is not None and e is not None and p is not None
    bu = str(bu).strip()
    want_orgs = None
    if orgs:
        seq = orgs.split(',') if isinstance(orgs, str) else orgs
        want_orgs = {str(o).strip() for o in seq if o and str(o).strip()}
    root = _INVENTORY_ROOT.format(m=mock)
    with pyodbc.connect(conn_str) as conn:
        cur = conn.cursor()
        r_schema, r_cols = _object_meta(cur, db, root)
        if not r_cols:
            return {"ok": False, "error": f"{root} not found in {db}"}
        org_col = _resolve_col(r_cols, _INVENTORY_ORG_FIELD)
        item_col = resolve_link_column(r_cols, _INVENTORY_LINK)
        if not (org_col and item_col):
            return {"ok": False, "error": "Organization or Item column not found on Items"}
        r_fq = _qualified(db, r_schema, root)
        # This BU's inventory orgs (INV_016651 -> 016 via _bu_matches), optionally
        # narrowed to a requested subset so one BU can split into separate files.
        cur.execute(f"SELECT DISTINCT [{org_col}] FROM {r_fq}")
        orgs = [r[0] for r in cur.fetchall()
                if r[0] is not None and _bu_matches(bu, str(r[0]).strip(), '')
                and (want_orgs is None or str(r[0]).strip() in want_orgs)]
        if not orgs:
            return {"ok": True, "bu": bu, "root": root, "population": 0, "sample_size": 0,
                    "root_key": item_col, "per_org": [], "sheets": []}
        ph = ",".join("?" * len(orgs))
        cur.execute(f"SELECT COUNT(*) FROM {r_fq} WHERE [{org_col}] IN ({ph})", orgs)
        population = cur.fetchone()[0]
        sheets = []
        r_headers, r_rows, per_org = None, [], []
        if population and stratify:
            # Stratified: size + draw each organization independently, then union.
            for org in orgs:
                cur.execute(f"SELECT COUNT(*) FROM {r_fq} WHERE [{org_col}] = ?", (org,))
                org_pop = cur.fetchone()[0]
                if not org_pop:
                    continue
                n_org = _cochran_n(org_pop, z, e, p)
                cur.execute(f"SELECT TOP ({n_org}) * FROM {r_fq} WHERE [{org_col}] = ? ORDER BY NEWID()", (org,))
                r_headers = [d[0] for d in cur.description]
                r_rows.extend([_coerce(v) for v in row] for row in cur.fetchall())
                per_org.append({"org": str(org).strip(), "population": org_pop, "sampled": n_org})
        elif population:
            n = max(1, min(n_req, population))
            cur.execute(f"SELECT TOP ({n}) * FROM {r_fq} WHERE [{org_col}] IN ({ph}) ORDER BY NEWID()", orgs)
            r_headers = [d[0] for d in cur.description]
            r_rows = [[_coerce(v) for v in row] for row in cur.fetchall()]
        if r_rows:
            sheets.append({"label": _short_name(root), "table": root, "headers": r_headers,
                           "rows": r_rows, "population": population})
            ii = r_headers.index(item_col)
            keys = sorted({row[ii] for row in r_rows if row[ii] is not None}, key=lambda x: str(x))
            for label, stem in _INVENTORY_CHILDREN:
                child = stem.format(m=mock)
                c_schema, c_cols = _object_meta(cur, db, child)
                if not c_cols:
                    continue
                c_link = resolve_link_column(c_cols, _INVENTORY_LINK)
                if not c_link:
                    continue
                c_fq = _qualified(db, c_schema, child)
                # True population = all this BU's child rows (every item, not just sampled),
                # for the Relationships/Sizing "Population rows" column.
                cur.execute(f"SELECT COUNT(*) FROM {c_fq} WHERE [{c_link}] IN "
                            f"(SELECT [{item_col}] FROM {r_fq} WHERE [{org_col}] IN ({ph}))", orgs)
                c_pop = cur.fetchone()[0]
                c_headers, c_rows = list(c_cols), []
                for i in range(0, len(keys), IN_CHUNK):
                    batch = keys[i:i + IN_CHUNK]
                    cph = ",".join("?" * len(batch))
                    cur.execute(f"SELECT * FROM {c_fq} WHERE [{c_link}] IN ({cph})", batch)
                    c_headers = [d[0] for d in cur.description]
                    c_rows.extend([[_coerce(v) for v in row] for row in cur.fetchall()])
                sheets.append({"label": label, "table": child, "headers": c_headers,
                               "rows": c_rows, "population": c_pop})
    return {"ok": True, "bu": bu, "root": root, "population": population,
            "sample_size": (len(sheets[0]["rows"]) if sheets else 0),
            "root_key": item_col, "per_org": per_org, "sheets": sheets}


# Assets sample scope (Entity_Hierarchy.docx): Assets master + Asset Distribution child,
# linked on 'Tag Number'. Keyed by a clean 'BU' column — matched EXACTLY, so BU 122's
# per-office books ('122_AGU','122_ARE',...) each sample as their own unit (client's rule;
# Assets is the only entity kept per-office).
_ASSETS_ROOT = 'FIN_ASSETS_{m}_VW_CONVERTED_TBL'
_ASSETS_CHILDREN = [('Asset Distribution', 'FIN_ASSETS_DISTRIBUTION_{m}_VW_CONVERTED_TBL')]
_ASSETS_BU_FIELD = 'BU'
_ASSETS_BOOK_FIELD = 'ASSET_BOOK'
_ASSETS_LINK = 'Tag Number'


def sample_assets_by_bu(conn_str, book, bu, sample_size, mock='MOCK14', source_db=None):
    """Server-side Assets sampling for ONE office = one (ASSET_BOOK, BU) pair. Every asset
    book is a separate office (client rule): BU 122's offices sit in the BU field
    (122_AGU...) while other BUs' offices sit in ASSET_BOOK (MAB_071651/652/...), so
    matching both columns exactly isolates a single office. Samples n assets + their Asset
    Distribution rows (linked on Tag Number)."""
    db = source_db or SOURCE_DATABASE
    try:
        n_req = int(sample_size)
    except (TypeError, ValueError):
        return {"ok": False, "error": "sample_size must be a number"}
    bu = str(bu).strip()
    book = str(book).strip()
    root = _ASSETS_ROOT.format(m=mock)
    with pyodbc.connect(conn_str) as conn:
        cur = conn.cursor()
        r_schema, r_cols = _object_meta(cur, db, root)
        if not r_cols:
            return {"ok": False, "error": f"{root} not found in {db}"}
        bu_col = _resolve_col(r_cols, _ASSETS_BU_FIELD)
        book_col = _resolve_col(r_cols, _ASSETS_BOOK_FIELD)
        tag_col = resolve_link_column(r_cols, _ASSETS_LINK)
        if not (bu_col and book_col and tag_col):
            return {"ok": False, "error": "BU, ASSET_BOOK or Tag Number column not found on Assets"}
        r_fq = _qualified(db, r_schema, root)
        cond, args = f"WHERE [{book_col}] = ? AND [{bu_col}] = ?", (book, bu)
        cur.execute(f"SELECT COUNT(*) FROM {r_fq} {cond}", args)
        population = cur.fetchone()[0]
        n = max(1, min(n_req, population)) if population else 0
        sheets = []
        if n:
            cur.execute(f"SELECT TOP ({n}) * FROM {r_fq} {cond} ORDER BY NEWID()", args)
            r_headers = [d[0] for d in cur.description]
            r_rows = [[_coerce(v) for v in row] for row in cur.fetchall()]
            sheets.append({"label": _short_name(root), "table": root, "headers": r_headers,
                           "rows": r_rows, "population": population})
            ti = r_headers.index(tag_col)
            keys = sorted({row[ti] for row in r_rows if row[ti] is not None}, key=lambda x: str(x))
            for label, stem in _ASSETS_CHILDREN:
                child = stem.format(m=mock)
                c_schema, c_cols = _object_meta(cur, db, child)
                if not c_cols:
                    continue
                c_link = resolve_link_column(c_cols, _ASSETS_LINK)
                if not c_link:
                    continue
                c_fq = _qualified(db, c_schema, child)
                # True population = all this office's child rows (every asset, not just
                # sampled), for the Relationships/Sizing "Population rows" column.
                cur.execute(f"SELECT COUNT(*) FROM {c_fq} WHERE [{c_link}] IN "
                            f"(SELECT [{tag_col}] FROM {r_fq} {cond})", args)
                c_pop = cur.fetchone()[0]
                c_headers, c_rows = list(c_cols), []
                for i in range(0, len(keys), IN_CHUNK):
                    batch = keys[i:i + IN_CHUNK]
                    cph = ",".join("?" * len(batch))
                    cur.execute(f"SELECT * FROM {c_fq} WHERE [{c_link}] IN ({cph})", batch)
                    c_headers = [d[0] for d in cur.description]
                    c_rows.extend([[_coerce(v) for v in row] for row in cur.fetchall()])
                sheets.append({"label": label, "table": child, "headers": c_headers,
                               "rows": c_rows, "population": c_pop})
    return {"ok": True, "book": book, "bu": bu, "root": root, "population": population,
            "sample_size": (len(sheets[0]["rows"]) if sheets else 0),
            "root_key": tag_col, "sheets": sheets}


def assets_bu_units(conn_str, mock='MOCK14', source_db=None):
    """Assets sample units — one per OFFICE = one distinct (ASSET_BOOK, BU) pair, each with
    its population and a display label. A BU with a single book keeps its plain label (015,
    122_AGU); a BU with multiple books labels each office by its ASSET_BOOK (MAB_071652)."""
    db = source_db or SOURCE_DATABASE
    root = _ASSETS_ROOT.format(m=mock)
    with pyodbc.connect(conn_str) as conn:
        cur = conn.cursor()
        r_schema, r_cols = _object_meta(cur, db, root)
        if not r_cols:
            return {"ok": False, "error": f"{root} not found in {db}"}
        bu_col = _resolve_col(r_cols, _ASSETS_BU_FIELD)
        book_col = _resolve_col(r_cols, _ASSETS_BOOK_FIELD)
        if not (bu_col and book_col):
            return {"ok": False, "error": "BU or ASSET_BOOK column not found on Assets"}
        r_fq = _qualified(db, r_schema, root)
        cur.execute(f"SELECT [{book_col}],[{bu_col}],COUNT(*) FROM {r_fq} "
                    f"WHERE [{bu_col}] IS NOT NULL GROUP BY [{book_col}],[{bu_col}]")
        rows = [(str(r[0]).strip(), str(r[1]).strip(), int(r[2])) for r in cur.fetchall() if str(r[1]).strip()]
    books_per_bu = {}
    for book, bu, _n in rows:
        books_per_bu.setdefault(bu, set()).add(book)
    units = []
    for book, bu, n in rows:
        label = bu if len(books_per_bu.get(bu, ())) <= 1 else book
        units.append({"book": book, "bu": bu, "label": label, "population": n})
    return {"ok": True, "root": root, "units": units}


def generate_entity_files(conn_str, s3, bucket, mock, entity, subentity=None,
                          dry_run=False, source_db=None, actor="", bu_filter=None):
    # bu_filter: when set, only source[/BU] splits whose source OR BU value equals
    # it are generated — so one agency's files can be produced in a small, fast
    # call (no 400-file cap / Lambda timeout) instead of the whole entity at once.
    bu_filter = (str(bu_filter).strip() if bu_filter not in (None, "") else None)
    ledger_segment = (entity or '').strip().lower() in _LEDGER_SEGMENT_ENTITIES
    db = source_db or SOURCE_DATABASE
    plan_table = f"SETUP_CONVERSION_PLAN_{mock}"
    out_prefix = f"{GENERATED_FOLDER}{mock}/{_safe_name(entity)}/"
    planned, generated, missing, empties = [], [], [], []

    with pyodbc.connect(conn_str) as conn:
        cur = conn.cursor()
        cur.execute(f"SELECT COUNT(*) FROM [{db}].INFORMATION_SCHEMA.TABLES WHERE TABLE_NAME = ?", (plan_table,))
        if cur.fetchone()[0] == 0:
            return {"ok": False, "error": f"{plan_table} not found in {db}"}

        where = ("ISNULL([CONVERSION_TABLE_BU],'') <> '' AND ISNULL([ENRICHMENT_SYSTEM],'') <> 'Y' "
                 "AND ISNULL([CONVERSION_TABLE_SourceField],'') <> '' AND [Entity] = ?")
        params = [entity]
        # Sampling-wired entities: honour the plan's "On Conversion Plan" flag live,
        # so only Y tables are generated (and a later Y flip is picked up on its own).
        if (entity or '').strip().lower() in _ON_PLAN_ENTITIES:
            where += " AND ISNULL([On Conversion Plan],'') = 'Y'"
        if subentity:
            where += " AND [SubEntity] = ?"
            params.append(subentity)
        cur.execute(
            f"SELECT DISTINCT [CONVERSION_TABLE_BU],[CONVERSION_TABLE_SourceField],[CONVERSION_TABLE_BU_Field] "
            f"FROM [{db}].[dbo].[{plan_table}] WHERE {where} ORDER BY 1", params)
        plans = cur.fetchall()
        # All conversion tables across every entity — a link-only child's parent can live
        # in a DIFFERENT entity (Item Category -> Items), so the per-entity `plans` alone
        # can't resolve it. Used only for parent lookup in _linked_child_files.
        cur.execute(
            f"SELECT DISTINCT [CONVERSION_TABLE_BU],[CONVERSION_TABLE_SourceField],[CONVERSION_TABLE_BU_Field] "
            f"FROM [{db}].[dbo].[{plan_table}] WHERE ISNULL([CONVERSION_TABLE_BU],'') <> ''")
        all_plans = cur.fetchall()

        for (table, src_field, bu_field) in plans:
            table = (table or '').strip()
            src_field = (src_field or '').strip()
            bu_field = (bu_field or '').strip()
            if not table or not src_field:
                continue
            prefix = _table_prefix(table)
            _, cols = _object_meta(cur, db, table)
            if not cols:
                missing.append(table)
                empties.append(_empty_entry(table, prefix, bu_filter, "not_built"))
                continue
            src_col = _resolve_col(cols, src_field)
            if not src_col:
                # Link-only child (no agency column of its own): build it from its
                # parent's rows, inheriting the parent's source/BU, so the tables stay
                # unchanged and the client links it by key like a normal child.
                linked = _linked_child_files(cur, db, table, prefix, bu_filter, ledger_segment, all_plans)
                if linked is not None:
                    produced = 0
                    for e in linked:
                        produced += 1
                        if dry_run:
                            planned.append({"file": e["file"], "table": table, "source": e["source"], "bu": e["bu"], "rows": len(e["rows"])})
                            continue
                        if len(generated) >= GEN_MAX_FILES:
                            break
                        s3.put_object(Bucket=bucket, Key=out_prefix + e["file"],
                                      Body=_rows_to_xlsx_bytes(e["headers"], e["rows"]), ContentType=XLSX_CONTENT_TYPE)
                        generated.append({"file": e["file"], "key": out_prefix + e["file"], "table": table,
                                          "source": e["source"], "bu": e["bu"], "rows": len(e["rows"])})
                    if produced == 0:
                        empties.append(_empty_entry(table, prefix, bu_filter, "no_rows_for_bu" if bu_filter else "empty_table"))
                    continue
                missing.append(f"{table} (no source field '{src_field}')")
                empties.append(_empty_entry(table, prefix, bu_filter, "no_source_field"))
                continue
            bu_col = _resolve_col(cols, bu_field) if bu_field else None
            fq = f"[{db}].[dbo].[{table}]"

            # Distinct source[/BU] combinations (excluding the header-as-value quirk).
            if bu_col and bu_col.lower() != src_col.lower():
                cur.execute(f"SELECT DISTINCT [{src_col}],[{bu_col}] FROM {fq} WHERE [{src_col}] <> ?", (src_field,))
            else:
                bu_col = None
                cur.execute(f"SELECT DISTINCT [{src_col}] FROM {fq} WHERE [{src_col}] <> ?", (src_field,))
            combos = cur.fetchall()

            used_names = set()
            produced = 0  # files this table yields for the current run / BU filter
            for combo in combos:
                raw_source = combo[0]
                if raw_source is None:
                    continue  # NULL source: excluded by the enumeration and by the scripts
                source = str(raw_source).strip()
                has_bu = bool(bu_col) and len(combo) > 1 and combo[1] is not None
                raw_bu = combo[1] if has_bu else None
                bu = str(raw_bu).strip() if has_bu else ''
                if bu_filter and not _bu_matches(bu_filter, source, bu, ledger_segment):
                    continue  # per-BU generation: skip splits for other agencies
                produced += 1
                # Match on the EXACT stored value (like the scripts) so blank or
                # padded source / BU values are captured, not skipped or mismatched.
                # A value that is blank after trimming is named "NA" in the file.
                if has_bu:
                    fname = f"CV_{prefix}__{_safe_name(source)}_{_safe_name(bu)}.xlsx"
                    cond = f"WHERE [{src_col}] = ? AND [{bu_col}] = ?"
                    args = (raw_source, raw_bu)
                else:
                    fname = f"CV_{prefix}__{_safe_name(source)}.xlsx"
                    cond = f"WHERE [{src_col}] = ?"
                    args = (raw_source,)
                base = fname[:-5]
                k = 2
                while fname in used_names:
                    fname = f"{base}_{k}.xlsx"
                    k += 1
                used_names.add(fname)

                if dry_run:
                    cur.execute(f"SELECT COUNT(*) FROM {fq} {cond}", args)
                    planned.append({"file": fname, "table": table, "source": source, "bu": bu,
                                    "rows": cur.fetchone()[0]})
                    continue

                if len(generated) >= GEN_MAX_FILES:
                    cap = f"stopped at {GEN_MAX_FILES} files"
                    mkey = _write_gen_manifest(s3, bucket, mock, entity, subentity, db,
                                               actor, out_prefix, generated, missing, capped=cap,
                                               bu_filter=bu_filter, empties=empties)
                    return {"ok": True, "entity": entity, "mock": mock, "folder": out_prefix,
                            "generated": generated, "missing": missing, "empties": empties,
                            "capped": cap, "manifest_key": mkey}

                cur.execute(f"SELECT * FROM {fq} {cond}", args)
                headers = [d[0] for d in cur.description]
                rows = cur.fetchall()
                s3.put_object(Bucket=bucket, Key=out_prefix + fname,
                              Body=_rows_to_xlsx_bytes(headers, rows), ContentType=XLSX_CONTENT_TYPE)
                generated.append({"file": fname, "key": out_prefix + fname, "table": table,
                                  "source": source, "bu": bu, "rows": len(rows)})

            if produced == 0:  # table exists but yielded no file for this run/BU
                empties.append(_empty_entry(
                    table, prefix, bu_filter,
                    "no_rows_for_bu" if bu_filter else "empty_table"))

        # Confirmed-orphan source families (e.g. Supplier '911' → BU 045): emit the
        # source's master + children re-tagged to the real agency so a per-BU run
        # surfaces them and the client's per-source split names the file accordingly.
        for fam in _ORPHAN_SOURCE_FAMILIES:
            if fam['entity'].strip().lower() != (entity or '').strip().lower():
                continue
            if bu_filter and bu_filter != fam['real_bu']:
                continue  # only on the assigned agency's run (or a full generate)
            for e in _orphan_source_family_files(cur, db, plans, fam):
                if dry_run:
                    planned.append({"file": e["file"], "table": e["table"],
                                    "source": e["source"], "bu": e["bu"], "rows": len(e["rows"])})
                    continue
                if len(generated) >= GEN_MAX_FILES:
                    break
                s3.put_object(Bucket=bucket, Key=out_prefix + e["file"],
                              Body=_rows_to_xlsx_bytes(e["headers"], e["rows"]), ContentType=XLSX_CONTENT_TYPE)
                generated.append({"file": e["file"], "key": out_prefix + e["file"], "table": e["table"],
                                  "source": e["source"], "bu": e["bu"], "rows": len(e["rows"])})

    if dry_run:
        return {"ok": True, "entity": entity, "mock": mock, "dry_run": True,
                "folder": out_prefix, "planned": planned, "missing": missing,
                "empties": empties, "plan_count": len(plans)}
    mkey = _write_gen_manifest(s3, bucket, mock, entity, subentity, db, actor,
                               out_prefix, generated, missing, bu_filter=bu_filter,
                               empties=empties)
    return {"ok": True, "entity": entity, "mock": mock, "folder": out_prefix,
            "generated": generated, "missing": missing, "empties": empties,
            "manifest_key": mkey}
