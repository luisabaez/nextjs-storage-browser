"""
cleanse_log.py — the Data Cleanse Log: every validation raised on a mock
cycle's files, with its occurrence count, as the validation team reports it.

The team keeps one set of views per mock cycle, each with the cycle written
into its definition:

  by BU     HCM_DATA_CLEANSE_BY_BU_<MOCK>_VW, FSCM_DATA_CLEANSE_BY_BU_<MOCK>_VW
            (older FSCM cycles: DATA_CLEANSE_BY_BU_<MOCK>_VW)
  summary   DATA_CLEANSE_HCM_<MOCK>_VW, DATA_CLEANSE_FSCM_<MOCK>_VW

The mock cycle is a request parameter here: the view for that cycle is looked
up by name, and a cycle that has no view yet is answered by an equivalent
query over the rule catalog and the detail log, so a new cycle works before
its views are created. The log holds counts only.

Agency users see the rows of their parties (source + agency; a party with a
blank agency is the source-level view and sees the whole source); super users
and certification reviewers see everything. The screen also receives the
rules behind the codes it shows — message, severity and path forward — which
is the agency users' read-only view of the validation rules.

In a test target whose log has nothing for the cycle yet, the main database
is read instead, so the screens have something to show.
"""
import io
import uuid
from datetime import datetime
from decimal import Decimal

import boto3
import pyodbc
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import api_util
import authz
import validation_seed
from api_util import ApiError

ACTIONS = {"cleanse_log", "cleanse_log_export"}

DB = validation_seed.TARGET_DB
SOURCE = validation_seed.SOURCE_DB

PILLARS = ("HCM", "FSCM")
LAYOUTS = ("by_bu", "summary")
EXPORT_PREFIX = f"{api_util.PRIVATE_ROOT}/DataCleanseLog"
CATALOG_TABLE = "SETUP_ERROR_MESSAGES_SOURCE"
DETAIL_TABLE = "LOG_DATA_CLEANSE_DETAIL"

MAX_ROWS = 20000          # rows returned to the screen
MAX_BYTES = 4_500_000     # approximate response size; a function URL stops at 6 MB
MAX_RULE_BYTES = 1_000_000  # the rules travel in the same response as the rows
EXPORT_MAX_ROWS = 100000
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ── request ──────────────────────────────────────────────────────────────────

def parse_request(values):
    """Validated (mock, pillar, layout, email, db and the optional source /
    agency / bu filters) from query-string or body values."""
    mock = api_util.mock(values.get("mock"))
    pillar = (values.get("pillar") or "").strip().upper() or ("HCM" if api_util.is_hcm_mock(mock) else "FSCM")
    if pillar not in PILLARS:
        raise ApiError(f"Invalid pillar: {pillar!r} (expected HCM or FSCM)")
    layout = (values.get("layout") or "by_bu").strip().lower()
    if layout not in LAYOUTS:
        raise ApiError(f"Invalid layout: {layout!r} (expected by_bu or summary)")
    email = (values.get("email") or values.get("actor") or "").strip()
    if not email:
        raise ApiError("email is required")
    return {"mock": mock, "pillar": pillar, "layout": layout, "email": email,
            "db": (values.get("db") or "").strip().lower(),
            "source": (values.get("source") or "").strip().upper(),
            "agency": authz.agency_code(values.get("agency") or ""),
            "bu": (values.get("bu") or "").strip().upper()}


# ── which view ───────────────────────────────────────────────────────────────

def view_candidates(pillar, layout, mock):
    """The team's view names for a pillar / layout / mock, in the order to try."""
    if layout == "summary":
        return [f"DATA_CLEANSE_{pillar}_{mock}_VW"]
    if pillar == "HCM":
        return [f"HCM_DATA_CLEANSE_BY_BU_{mock}_VW"]
    return [f"FSCM_DATA_CLEANSE_BY_BU_{mock}_VW", f"DATA_CLEANSE_BY_BU_{mock}_VW"]


def _find_view(cur, db, candidates):
    """The first candidate that exists in `db`, spelled as the database has it
    (the team's names vary in case, e.g. ..._Mock12_VW)."""
    for name in candidates:
        cur.execute(f"SELECT name FROM [{db}].sys.views WHERE UPPER(name) = ?", (name.upper(),))
        row = cur.fetchone()
        if row:
            return api_util.ident(row[0], "view name")
    return None


def _exists(cur, db, name):
    cur.execute(f"SELECT COUNT(*) FROM [{db}].sys.objects WHERE name = ?", (name,))
    return cur.fetchone()[0] > 0


def _clone(conn, names, warnings):
    """Test target only: create `names` there from the source's definitions."""
    try:
        seeder = validation_seed.Seeder(conn)
        for name in names:
            seeder.ensure(name)
        warnings.extend(f"Could not copy {failure}" for failure in seeder.report["failed"])
    except Exception as e:  # noqa: BLE001 - reported to the caller as a warning
        warnings.append(f"Could not prepare {DB}: {str(e)[:200]}")


def _resolve_view(conn, cur, candidates, read_db, warnings):
    """Name of the view to read in `read_db`, or None when the mock has none.
    A view the test target lacks is copied from the source database first."""
    view = _find_view(cur, read_db, candidates)
    if view or read_db.upper() == SOURCE.upper():
        return view
    source_view = _find_view(cur, SOURCE, candidates)
    if not source_view:
        return None
    _clone(conn, [source_view], warnings)
    view = _find_view(cur, read_db, candidates)
    if view:
        warnings.append(f"{view} was copied from {SOURCE} into {read_db}")
    else:
        warnings.append(f"{source_view} exists in {SOURCE} but could not be created in {read_db}")
    return view


# ── a cycle without views ────────────────────────────────────────────────────

def fallback_sql(db, mock, pillar, layout):
    """(sql, args) equivalent to the team's views for a mock that has none yet:
    the rule catalog joined to the detail log's counts for the mock. `db` is one
    of this module's two database names and `mock` comes from api_util.mock().
    Headers follow the views, [Error Messge] included, so workbooks line up."""
    by_bu = layout == "by_bu"
    keys = "[Validation_Code], [Validation_Program], [Source]" + (", [BU]" if by_bu else "")
    detail = (f"SELECT {keys}, COUNT(*) AS [Cnt], MAX([Validation_PROCESSED_DTTM]) AS [LastRun] "
              f"FROM [{db}].dbo.[{DETAIL_TABLE}] WHERE [MOCK] = ? GROUP BY {keys}")
    octane = "e.[OctaneA], e.[OctaneB], e.[OctaneC], e.[OctaneD], "
    if by_bu:
        select = (
            "e.[Pillar] AS [Pillar], e.[VALIDATION_CODE] AS [Validation Code], e.[ERROR_MESSAGE] AS [Error Messge], "
            "e.[VALIDATION_TYPE] AS [Validation Type], e.[ENTITY] AS [Entity], d.[BU] AS [BU], "
            "e.[AgencyReports] AS [Agency Reports], e.[SourceReports] AS [Source Reports], d.[Source] AS [Source], "
            f"d.[Cnt] AS [{mock} Count], {octane}e.[Severity], e.[Severity_Criteria], "
            "e.[TransformationLogicApplied], e.[NotInScope], e.[RecordsConverted], e.[Notes], "
            "d.[Validation_Program] AS [Validation_Program]")
        reported = " AND (e.[AgencyReports] = 'Y' OR e.[SourceReports] = 'Y')"
        order = "e.[VALIDATION_CODE], d.[Source], d.[BU]"
    else:
        select = (
            "e.[Pillar] AS [Pillar], d.[Validation_Program] AS [Validation Program], "
            "e.[VALIDATION_CODE] AS [Validation Code], e.[ERROR_MESSAGE] AS [Error Messge], "
            "e.[VALIDATION_TYPE] AS [Validation Type], e.[ENTITY] AS [Entity], d.[Source] AS [Source], "
            f"d.[Cnt] AS [{mock} Count], d.[LastRun] AS [Date], {octane}e.[Severity], "
            "e.[AgencyReports] AS [Reported to Agencies], e.[SourceReports] AS [Reported to Sources]")
        reported = ""   # the summary shows the two flags as columns instead of filtering on them
        order = "e.[VALIDATION_CODE], d.[Source]"
    # Every rule that is not HCM belongs to the FSCM log.
    pillar_clause = "e.[Pillar] = 'HCM'" if pillar == "HCM" else "ISNULL(e.[Pillar], '') <> 'HCM'"
    sql = (f"SELECT {select} FROM [{db}].dbo.[{CATALOG_TABLE}] e "
           f"JOIN ({detail}) d ON d.[Validation_Code] = e.[VALIDATION_CODE] "
           f"WHERE {pillar_clause}{reported} AND ISNULL(d.[Source], '') NOT LIKE '%BEFOREORACLE%' "
           f"ORDER BY {order}")
    return sql, (mock,)


# ── reading ──────────────────────────────────────────────────────────────────

def _cell(v):
    """A JSON / Excel friendly value."""
    if v is None or isinstance(v, (int, float, str)):
        return v
    if isinstance(v, Decimal):
        return int(v) if v == v.to_integral_value() else float(v)
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    return str(v)


def _index(columns, name):
    lowered = [c.strip().lower() for c in columns]
    return lowered.index(name) if name in lowered else None


def _limited_to_agencies(email):
    """Whether an agency user's rights name agencies within a source: parties
    with an agency, or the older business-unit list."""
    parties = authz.parties(email)
    if parties:
        return any(agency for _, agency in parties)
    return bool(authz.get_permissions(email).get("allowedBusinessUnits"))


def row_filters(req, restricted):
    """(allowed, wanted): two tests of a row's (source, agency, bu). `allowed`
    is the caller's parties; `wanted` is the request's own filters. The log has
    no Agency column, so a row's agency is read from its BU, whose first three
    characters are the agency number (the same reading authz makes)."""
    email, decided = req["email"], {}

    def allowed(source, agency, bu):
        if not restricted:
            return True
        key = (source, agency, bu)
        if key not in decided:   # thousands of rows, a handful of distinct parties
            # A party with a blank agency is the source-level view: the whole source.
            decided[key] = (authz.can_act_on_party(email, source, agency, bu)
                            or authz.can_act_on_party(email, source, ""))
        return decided[key]

    def wanted(source, agency, bu):
        return ((not req["source"] or source.upper() == req["source"])
                and (not req["agency"] or req["agency"] in (authz.agency_code(agency), bu[:3].upper()))
                and (not req["bu"] or bu.upper() == req["bu"]))

    return allowed, wanted


def read_rows(cur, allowed, wanted, max_rows, max_bytes=None):
    """(columns, rows, total, sources, seen) of the executed cursor. `seen`
    counts what the database returned; `sources` are those of the rows the
    caller may see, so the screen's source list survives a source filter;
    `rows` / `total` are the allowed rows the request asked for."""
    columns = [d[0] for d in cur.description]
    i_source, i_agency, i_bu = (_index(columns, name) for name in ("source", "agency", "bu"))

    def text(r, i):
        return str(r[i] or "").strip() if i is not None else ""

    rows, sources, seen, total, size = [], set(), 0, 0, 0
    for r in cur:
        seen += 1
        source, bu = text(r, i_source), text(r, i_bu)
        agency = text(r, i_agency) or bu
        if not allowed(source, agency, bu):
            continue
        if source:
            sources.add(source)
        if not wanted(source, agency, bu):
            continue
        total += 1
        if len(rows) < max_rows and (max_bytes is None or size < max_bytes):
            row = [_cell(v) for v in r]
            size += sum(len(str(v)) for v in row if v is not None) + 5 * len(row)
            rows.append(row)
    return columns, rows, total, sorted(sources), seen


def _read(conn, cur, req, layout, read_db, filters, limits, warnings):
    """(view name or None, read_rows result) for the log as `read_db` holds it."""
    mock, pillar = req["mock"], req["pillar"]
    view = _resolve_view(conn, cur, view_candidates(pillar, layout, mock), read_db, warnings)
    if view:
        sql, args = f"SELECT * FROM [{read_db}].dbo.[{view}]", ()
    else:
        if read_db.upper() != SOURCE.upper():
            missing = [t for t in (CATALOG_TABLE, DETAIL_TABLE) if not _exists(cur, read_db, t)]
            if missing:
                _clone(conn, missing, warnings)
        sql, args = fallback_sql(read_db, mock, pillar, layout)
        warnings.append(f"No {pillar} Data Cleanse Log view exists for {mock}; "
                        f"the log was built from {CATALOG_TABLE} and {DETAIL_TABLE}")
    cur.execute(sql, args)
    return view, read_rows(cur, *filters, *limits)


def _catalog_columns(cur):
    cur.execute(f"SELECT COLUMN_NAME FROM [{DB}].INFORMATION_SCHEMA.COLUMNS WHERE TABLE_NAME = ?", (CATALOG_TABLE,))
    return {r[0].upper() for r in cur.fetchall()}


def read_rules(conn, cur, columns, rows, warnings):
    """{code: rule} for the validation codes in `rows`: what each rule says and
    its path forward, from the rule catalog the application edits. The team's
    Notes / InternalNote are never part of it."""
    i_code = _index(columns, "validation code")
    if i_code is None:
        return {}
    # Upper-cased for the lookup, kept as the rows spell it for the answer.
    codes = {str(r[i_code]).strip().upper(): str(r[i_code]).strip() for r in rows if r[i_code]}
    if not codes:
        return {}
    have = _catalog_columns(cur)
    if not have and DB.upper() != SOURCE.upper():
        _clone(conn, [CATALOG_TABLE], warnings)
        have = _catalog_columns(cur)
    if not have:
        warnings.append(f"{CATALOG_TABLE} does not exist in {DB}; the rules are not included")
        return {}
    # Two columns the team added later; an older copy of the table lacks them.
    spanish, forward = (f"[{c}]" if c in have else f"NULL AS [{c}]" for c in ("ERROR_MESSAGE_SPA", "PATH_FORWARD"))
    rules, wanted, size = {}, sorted(codes), 0
    for start in range(0, len(wanted), 500):   # SQL Server takes at most 2,100 parameters
        batch = wanted[start:start + 500]
        cur.execute(
            f"SELECT [VALIDATION_CODE], [ERROR_MESSAGE], {spanish}, {forward}, [Severity], [ENTITY], [VALIDATION_TYPE] "
            f"FROM [{DB}].dbo.[{CATALOG_TABLE}] WHERE UPPER(LTRIM(RTRIM([VALIDATION_CODE]))) IN "
            f"({', '.join('?' * len(batch))})", batch)
        for r in cur.fetchall():
            # The database compares by its collation, so a match need not be one here.
            code = codes.get(str(r[0]).strip().upper())
            if code is None or code in rules:
                continue
            size += 100 + sum(len(str(v)) for v in r if v is not None)
            if size > MAX_RULE_BYTES:
                warnings.append(f"Only the rules of the first {len(rules):,} validation codes are included")
                return rules
            rules[code] = {"message": r[1], "message_spa": r[2], "path_forward": r[3],
                           "severity": r[4], "entity": r[5], "type": r[6]}
    return rules


def load(conn_str, req, max_rows, max_bytes=None, with_rules=False):
    """The log for a parsed request, from the mock's view or the equivalent query."""
    mock, pillar, layout = req["mock"], req["pillar"], req["layout"]
    role = authz.require(req["email"], authz.AGENCY_USER, authz.CERT_REVIEWER, what="viewing the Data Cleanse Log")
    restricted = role == authz.AGENCY_USER
    is_test = validation_seed.is_test_target()
    read_db = SOURCE if (is_test and req["db"] == "main") else DB
    warnings = []
    if layout == "summary" and (req["agency"] or req["bu"] or (restricted and _limited_to_agencies(req["email"]))):
        # The summary has no BU column, so an agency's part of it cannot be told apart.
        layout = "by_bu"
        warnings.append("The summary is not kept by agency; the log by BU is shown instead")
    filters, limits = row_filters(req, restricted), (max_rows, max_bytes)
    with pyodbc.connect(conn_str, autocommit=False) as conn:
        cur = conn.cursor()
        read_warnings = []
        view, result = _read(conn, cur, req, layout, read_db, filters, limits, read_warnings)
        if is_test and req["db"] != "main" and result[4] == 0:
            # Nothing has been run in the test target for this cycle yet.
            read_db = SOURCE
            read_warnings = [f"The test database has no log rows for {mock}; showing the main database."]
            view, result = _read(conn, cur, req, layout, read_db, filters, limits, read_warnings)
        warnings += read_warnings
        columns, rows, total, sources, _ = result
        rules = read_rules(conn, cur, columns, rows, warnings) if with_rules else None
    if restricted:
        # The team's working notes are not for the agencies.
        keep = [i for i, c in enumerate(columns) if c.strip().lower() not in ("notes", "internalnote")]
        columns, rows = [columns[i] for i in keep], [[r[i] for i in keep] for r in rows]
    if len(rows) < total:
        warnings.append(f"Only the first {len(rows):,} of {total:,} rows are included")
    data = {"mock": mock, "pillar": pillar, "layout": layout, "db": read_db, "is_test": is_test,
            "source": "view" if view else "query", "view_name": view,
            "columns": columns, "rows": rows, "total": total, "truncated": len(rows) < total,
            "sources": sources, "warnings": warnings}
    if with_rules:
        data["rules"] = rules
    return data


# ── export ───────────────────────────────────────────────────────────────────

def export_name(pillar, layout, mock, when=None):
    """The team's workbook naming."""
    stamp = (when or datetime.now()).strftime("%Y.%m.%d-%H.%M.%S")
    if layout == "summary":
        return f"DATA_CLEANSE_{pillar}_{mock}_{stamp}.xlsx"
    return f"{pillar}_DATA_CLEANSE_BY_BU_{mock}_{stamp}.xlsx"


def build_workbook(title, columns, rows):
    """Workbook bytes: one sheet, the log as returned."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(list(columns))
    fill, font = PatternFill("solid", fgColor="1F4E79"), Font(bold=True, color="FFFFFF")
    for c in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = fill
        cell.font = font
    for row in rows:
        ws.append([api_util.xlsx_value(ws, v) for v in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max(1, len(columns)))}{max(2, len(rows) + 1)}"
    for i, col in enumerate(columns):
        longest = max([len(str(col))] + [len(str(r[i])) for r in rows[:500] if r[i] is not None])
        ws.column_dimensions[get_column_letter(i + 1)].width = min(60, max(10, longest + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _export(conn_str, req, bucket):
    data = load(conn_str, req, EXPORT_MAX_ROWS)
    name = export_name(req["pillar"], data["layout"], req["mock"])
    # Unique per request: an export is filtered to its caller, so two
    # requests must never share one key. The download keeps the team's name.
    key = f"{EXPORT_PREFIX}/{req['mock']}/{name[:-len('.xlsx')]}_{uuid.uuid4().hex[:8]}.xlsx"
    content = build_workbook(name[:-len(".xlsx")], data["columns"], data["rows"])
    boto3.client("s3").put_object(Bucket=bucket, Key=key, Body=content, ContentType=_XLSX)
    return {"key": key, "name": name, "rows": len(data["rows"]), "truncated": data["truncated"],
            "warnings": data["warnings"], "url": api_util.presign_get(bucket, key, name)}


def handle(action, event, bucket, headers, conn_str):
    def run():
        if action == "cleanse_log":
            req = parse_request(api_util.params(event))
            return api_util.ok(headers, load(conn_str, req, MAX_ROWS, MAX_BYTES, with_rules=True))
        req = parse_request(api_util.body(event))
        return api_util.ok(headers, _export(conn_str, req, bucket))

    return api_util.guarded(run, headers)
