"""
cert_forms.py — the validation team's signed certification form.

Agencies certify on the team's own Excel form: one row per data entity with
the agency resource and one of the form's comments, then Name, Title, Date and
a signature. The portal hands out the cycle's blank form with the agency filled
in, takes the signed Excel back, keeps it for later review and reads its rows.

Templates: a super user uploads the cycle's blank forms once, one per kind:

  HR        agencies that certify only validations, HR module
  PAYROLL   the same agencies, Payroll and Compensation module
  SOURCES   the sources (and agencies outside RHUM): validations, converted
            data and pre-load recon on one form

stored at SymphonyPrivate/CertForms/<MOCK>/templates/<KIND>.xlsx. Which forms a
party needs follows certifications.form_kind.

Reading a signed form: on its Certification sheet, the table under the header
Segment / Data Entity / Folder (or Archivo) / Agency Resource / Comments, down
to the first row with no data entity, and the Agency, Name, Title and Date
cells next to their labels. A comment counts when it is one of the form's
options (accents, case and spacing ignored).

A form certifies the party's records it covers. A record takes the answer of
the form rows of its folder and segment: the row for the same data entity
when there is one, otherwise the most serious answer among those rows (issues
before agreement before no errors). Every row answered as incorrect becomes an
issue on its record, and needs a supporting document before the party signs
off (certifications._signoff).

Signed PDF: an agency that signs the form on paper (or prints it to PDF and
signs that) uploads the PDF. A PDF cannot be read, so the uploader records the
form's answers in the portal, row by row as on the form; the PDF is kept as
the signed form and the answers are recorded like an uploaded form's.

Electronic signature: instead of downloading, signing and uploading, the
agency answers the same rows in the portal and signs with its name and title.
The portal writes the answers and a signature block (name, e-mail of the
signed-in account, date and time) into the cycle's form, keeps that file and
reads it exactly like an uploaded one. Until the validation team approves it,
only super users can sign this way (setting esign_agencies turns it on for
agencies).

Tables (application database):
  DATA_CLEANSE_CERT_FORM       one uploaded or signed form: party, kind, file, what was read
  DATA_CLEANSE_CERT_FORM_ROW   its rows
Certifications and issues it produced carry its ID (Form_ID).
"""
import hashlib
import io
import re
import unicodedata
import uuid
import zipfile
from datetime import date, datetime, timedelta, timezone
from xml.sax.saxutils import escape

import boto3
import openpyxl
import pyodbc

import api_util
import app_settings
import authz
import certifications as certs
from api_util import ApiError

ACTIONS = {
    "certform_templates", "certform_template_upload_url", "certform_template_url", "certform_template_delete",
    "certform_status", "certform_list", "certform_rows", "certform_file_url",
    "certform_download", "certform_upload_url", "certform_submit", "certform_revoke",
    "certform_esign_form", "certform_esign", "certform_esign_setting",
}
_READS = {"certform_templates", "certform_template_url", "certform_status", "certform_list", "certform_rows",
          "certform_file_url", "certform_esign_form"}

DB = certs.DB
T_FORM = "DATA_CLEANSE_CERT_FORM"
T_ROW = "DATA_CLEANSE_CERT_FORM_ROW"

FORMS_PREFIX = f"{api_util.PRIVATE_ROOT}/CertForms"
KINDS = {"HR": "HCM-HR", "PAYROLL": "HCM-Payroll", "SOURCES": "Sources"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF = "application/pdf"
MAX_FORM_BYTES = 25 * 1024 * 1024
PDF_STATEMENT = "The answers recorded in the portal are those of the signed PDF form."

ESIGN_SETTING = "esign_agencies"   # "on" once electronic signature is approved for agencies
ESIGN_CONSENT = ("Firmo este formulario electrónicamente con mi nombre, título y cuenta del portal. Acepto que esta "
                 "firma electrónica tiene la misma validez que mi firma manuscrita.")
_PR_TIME = timezone(timedelta(hours=-4), "AST")   # Puerto Rico does not change the clock

_DDL = {
    T_FORM: (
        "[ID] INT IDENTITY(1,1) PRIMARY KEY, [MOCK] VARCHAR(20) NOT NULL, [Source] VARCHAR(50) NOT NULL, "
        "[Agency] VARCHAR(100) NOT NULL, [BU] VARCHAR(100) NOT NULL, [Kind] VARCHAR(20) NOT NULL, "
        "[File_Name] NVARCHAR(300) NOT NULL, [S3_Key] NVARCHAR(600) NOT NULL, [Size_Bytes] BIGINT NULL, "
        "[Form_Agency] NVARCHAR(200) NULL, [Signer_Name] NVARCHAR(200) NULL, [Signer_Title] NVARCHAR(200) NULL, "
        "[Signed_Date] NVARCHAR(50) NULL, [Has_Signature_Image] BIT NULL, [Rows_Total] INT NULL, "
        "[Rows_Answered] INT NULL, [Uploaded_By] NVARCHAR(200) NULL, [Uploaded_DTTM] DATETIME NULL, "
        "[Is_Current] BIT NOT NULL DEFAULT 1, [Revoked_By] NVARCHAR(200) NULL, [Revoked_DTTM] DATETIME NULL, "
        "[Revoke_Reason] NVARCHAR(MAX) NULL"),  # plus _ESIGN_COLUMNS
    T_ROW: (
        "[ID] INT IDENTITY(1,1) PRIMARY KEY, [Form_ID] INT NOT NULL, [Row_Number] INT NULL, "
        "[Segment] NVARCHAR(100) NULL, [Data_Entity] NVARCHAR(200) NULL, [Folder] NVARCHAR(100) NULL, "
        "[Agency_Resource] NVARCHAR(200) NULL, [Comment] NVARCHAR(1000) NULL, [Response_Code] VARCHAR(20) NULL"),
}

# How serious an answer is: a record answered on several rows takes the most serious.
_WEIGHT = {"ISSUES": 3, "AGREE": 2, "AGREE_EXCLUSIONS": 2, "COST_ALLOCATION": 2,
           "NO_ERRORS": 1, "NO_EXCLUSIONS": 1, "CONVERTED_OK": 1}
# How a form was signed: Sign_Method ESIGN for one signed in the portal, PDF for
# a signed PDF whose answers were recorded in the portal (NULL: the Excel form
# uploaded), with where it came from and a fingerprint of the file.
_ESIGN_COLUMNS = (("Sign_Method", "VARCHAR(20)"), ("Signer_IP", "VARCHAR(64)"), ("Signer_Agent", "NVARCHAR(400)"),
                  ("Content_SHA256", "CHAR(64)"), ("Consent_Text", "NVARCHAR(1000)"))

_HEADERS = {"segment": "segment", "data entity": "entity", "folder": "folder", "archivo": "folder",
            "agency resource": "resource", "comments": "comment"}
_LABELS = {"agency": "agency", "name": "name", "title": "title", "date": "date"}
_SOURCE_LINE = re.compile(r"^(Certifico que el sistema fuente de nuestros datos es )(.+?)(\.?)$", re.I)


# ── small pure helpers ───────────────────────────────────────────────────────

_s = certs._s


def _plain(text):
    """Text for comparing: no accents, no case, no punctuation, single spaces."""
    t = unicodedata.normalize("NFKD", _s(text))
    t = "".join(ch for ch in t if not unicodedata.combining(ch)).lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", t).split())


_COMMENTS = {_plain(text): code for code, text in certs.RESPONSES.items()}


def match_comment(comment):
    """The response code of one of the form's comments, or None."""
    return _COMMENTS.get(_plain(comment))


def _singular(word):
    if word.endswith(("xes", "ses", "zes")):
        return word[:-2]
    return word[:-1] if word.endswith("s") and not word.endswith("ss") and len(word) > 3 else word


def _entity_key(name):
    """A data entity as both the form and the setup table can spell it:
    "Person NID" / "Person National Identifier", "Departments" / "Department",
    "Assignment" / "Person Assignment"."""
    words = " ".join("national identifier" if w == "nid" else w for w in _plain(name).split()).split()
    if len(words) > 1 and words[0] == "person":
        words = words[1:]
    return " ".join(_singular(w) for w in words)


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    return _s(value)


def overall(codes):
    """The answer a record takes from several rows: the most serious, and among
    those the most frequent."""
    top = max(_WEIGHT[c] for c in codes)
    tied = [c for c in codes if _WEIGHT[c] == top]
    return max(sorted(set(tied)), key=tied.count)


def agency_label(party):
    """How the form's agency list names a party: "018 Junta De Planificacion";
    a source-level party is its source."""
    if not _s(party["agency"]):
        return _s(party["source"])
    bu = _s(party.get("bu"))
    m = re.match(r"^(\d{3})\s*[-–]?\s*(.*)$", bu)
    return f"{m.group(1)} {m.group(2)}".strip() if m else (bu or _s(party["agency"]))


def template_key(mock, kind):
    return f"{FORMS_PREFIX}/{mock}/templates/{kind}.xlsx"


def download_name(mock, kind, party):
    label = api_util.safe_segment(agency_label(party).replace(" ", "_"), "party")
    return f"{mock}_{KINDS[kind].replace(' ', '_')}_Data_Validation_Certification_{label}.xlsx"


# ── reading a filled-in form ─────────────────────────────────────────────────

def _cert_sheet(wb):
    for ws in wb.worksheets:
        if ws.title.strip().lower() == "certification":
            return ws
    return next((ws for ws in wb.worksheets if ws.sheet_state == "visible"), wb.worksheets[0])


def _label_cells(ws, max_row=None):
    """{label: (row, column)} of the Agency / Name / Title / Date labels."""
    found = {}
    for row in ws.iter_rows(max_row=max_row):
        for c in row:
            key = _plain(c.value) if isinstance(c.value, str) else ""
            if key in _LABELS and _s(c.value).endswith(":") and _LABELS[key] not in found:
                found[_LABELS[key]] = (c.row, c.column)
    return found


def _right_of(ws, row, col):
    """The first filled cell right of a label (the answer cells are merged)."""
    for c in range(col + 1, col + 7):
        text = _cell_text(ws.cell(row, c).value)
        if text:
            return text
    return ""


def _table_header(ws):
    """(header row, {field: column}) of the form's table."""
    for row in ws.iter_rows(max_row=40):
        found = {_HEADERS[_plain(c.value)]: c.column for c in row
                 if isinstance(c.value, str) and _plain(c.value) in _HEADERS}
        if {"entity", "comment"} <= set(found):
            return row[0].row, found
    raise ApiError("This file is not the certification form: the table header (Segment, Data Entity, "
                   "Folder, Agency Resource, Comments) was not found")


def read_form(content):
    """What a filled-in form says: {agency, name, title, date, signature_image,
    rows:[{row, segment, entity, folder, resource, comment, code}]}."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        signature = any(n.startswith("xl/media/") for n in zipfile.ZipFile(io.BytesIO(content)).namelist())
    except Exception:  # noqa: BLE001 - anything unreadable is simply not the form
        raise ApiError("The file could not be read as an Excel workbook (.xlsx)")
    ws = _cert_sheet(wb)
    header, cols = _table_header(ws)

    def get(r, field):
        return _cell_text(ws.cell(r, cols[field]).value) if field in cols else ""

    rows, r = [], header + 1
    while r <= ws.max_row and get(r, "entity"):
        comment = get(r, "comment")
        rows.append({"row": r, "segment": get(r, "segment"), "entity": get(r, "entity"),
                     "folder": get(r, "folder"), "resource": get(r, "resource"), "comment": comment,
                     "code": match_comment(comment) if comment else None})
        r += 1
    labels = {k: _right_of(ws, *at) for k, at in _label_cells(ws).items()}
    return {"agency": labels.get("agency", ""), "name": labels.get("name", ""), "title": labels.get("title", ""),
            "date": labels.get("date", ""), "signature_image": signature, "rows": rows}


def form_problems(form, party):
    """(errors that refuse the upload, warnings shown with it)."""
    errors, warnings = [], []
    if not form["rows"]:
        errors.append("The form has no data entity rows")
    answered = [x for x in form["rows"] if x["code"]]
    if form["rows"] and not answered:
        errors.append("No row has one of the form's comments")
    if not form["name"]:
        errors.append("Fill in Name in the signature block")
    # Only an agency number can tell a form for another agency apart: the
    # Sources form names the source in whatever spelling its list offers.
    wanted, given = agency_label(party), form["agency"]
    if re.match(r"^\d{3}$", authz.agency_code(given)) and authz.agency_code(given) != authz.agency_code(wanted):
        errors.append(f"This form is for {given}, not {wanted}")
    unknown = [x for x in form["rows"] if x["comment"] and not x["code"]]
    if unknown:
        warnings.append("These rows have a comment that is not one of the form's options and were not counted: "
                        + ", ".join(x["entity"] for x in unknown[:10]))
    blank = [x for x in form["rows"] if not x["comment"]]
    if blank:
        warnings.append(f"{len(blank)} row(s) have no comment: " + ", ".join(x["entity"] for x in blank[:10])
                        + (" …" if len(blank) > 10 else ""))
    if not form["title"] or not form["date"]:
        warnings.append("Title or Date is empty in the signature block")
    if not form["signature_image"] and not form.get("electronic") and not form.get("pdf"):
        warnings.append("No signature image was found in the file")
    return errors, warnings


def assign(records, rows):
    """{record key: (response code, rows)} for the records the answered rows
    cover (see the module notes)."""
    answered = [x for x in rows if x["code"]]
    out = {}
    for rec in records:
        folder, module = certs.file_class(rec["file_type"]), certs.module_class(rec["module"])
        pool = [x for x in answered if certs.file_class(x["folder"] or "Validation") == folder
                and (not module or not certs.module_class(x["segment"])
                     or certs.module_class(x["segment"]) == module)]
        same = [x for x in pool if _entity_key(x["entity"]) == _entity_key(rec["entity"])]
        chosen = same or pool
        if chosen:
            out[certs._row_key(rec)] = (overall([x["code"] for x in chosen]), chosen)
    return out


# ── filling a template ───────────────────────────────────────────────────────

def _sheet_path(z, title):
    """The worksheet part of the sheet titled `title`."""
    workbook = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    m = re.search(rf'<sheet\b[^>]*\bname="{re.escape(escape(title, {chr(34): "&quot;"}))}"[^>]*\br:id="([^"]+)"',
                  workbook)
    if not m:
        m = re.search(r'<sheet\b[^>]*\br:id="([^"]+)"', workbook)
    target = re.search(rf'<Relationship\b[^>]*\bId="{m.group(1)}"[^>]*\bTarget="([^"]+)"', rels)
    if not target:
        target = re.search(rf'<Relationship\b[^>]*\bTarget="([^"]+)"[^>]*\bId="{m.group(1)}"', rels)
    path = target.group(1).lstrip("/")
    return path if path.startswith("xl/") else f"xl/{path}"


def _column_number(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def set_cell(xml, ref, text):
    """The worksheet XML with cell `ref` holding `text` (an inline string that
    keeps the cell's style)."""
    new = f'<t xml:space="preserve">{escape(text)}</t>'
    cell = re.compile(rf'<c r="{ref}"(?P<attrs>[^>]*?)(?:/>|>.*?</c>)', re.S)
    m = cell.search(xml)
    if m:
        style = re.search(r'\ss="(\d+)"', m.group("attrs"))
        style = f' s="{style.group(1)}"' if style else ""
        return xml[:m.start()] + f'<c r="{ref}"{style} t="inlineStr"><is>{new}</is></c>' + xml[m.end():]
    col, row = re.match(r"([A-Z]+)(\d+)", ref).groups()
    row_m = re.search(rf'<row r="{row}"[^>]*?(/>|>)', xml)
    if not row_m:
        return xml
    element = f'<c r="{ref}" t="inlineStr"><is>{new}</is></c>'
    if row_m.group(1) == "/>":
        opened = row_m.group(0)[:-2] + ">"
        return xml[:row_m.start()] + opened + element + "</row>" + xml[row_m.end():]
    end = xml.index("</row>", row_m.end())
    at = end
    for c in re.finditer(r'<c r="([A-Z]+)\d+"', xml[row_m.end():end]):
        if _column_number(c.group(1)) > _column_number(col):
            at = row_m.end() + c.start()
            break
    return xml[:at] + element + xml[at:]


def fill_template(content, party):
    """The blank form with the party's agency and source written in. Edited in
    the file's XML: a round trip through openpyxl would drop the form's
    dropdown lists."""
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = _cert_sheet(wb)
    edits = {}
    labels = _label_cells(ws, max_row=20)
    if "agency" in labels:
        row, col = labels["agency"]
        edits[ws.cell(row, col + 1).coordinate] = agency_label(party)
    for row in ws.iter_rows():
        for c in row:
            m = _SOURCE_LINE.match(_s(c.value)) if isinstance(c.value, str) else None
            if m:
                edits[c.coordinate] = f"{m.group(1)}{_s(party['source'])}{m.group(3)}"
    return _write_cells(content, ws.title, edits)


def sign_form(content, answers, signer):
    """A filled-in form (from fill_template) with the portal's answers and an
    electronic signature block. answers: {row number: (resource, comment)};
    signer: {name, title, email, at}."""
    ws = _cert_sheet(openpyxl.load_workbook(io.BytesIO(content)))
    _, cols = _table_header(ws)
    edits = {}
    for row, (resource, comment) in answers.items():
        if resource and "resource" in cols:
            edits[ws.cell(row, cols["resource"]).coordinate] = resource
        if comment:
            edits[ws.cell(row, cols["comment"]).coordinate] = comment
    labels = _label_cells(ws)
    at = signer["at"]
    for key, text in (("name", signer["name"]), ("title", signer["title"]), ("date", at.strftime("%Y-%m-%d"))):
        if key in labels:
            row, col = labels[key]
            edits[ws.cell(row, col + 1).coordinate] = text
    signature = next((c for r in ws.iter_rows() for c in r
                      if isinstance(c.value, str) and _plain(c.value) == "signature"), None)
    if signature:
        edits[ws.cell(signature.row, signature.column + 1).coordinate] = (
            f"Firmado electrónicamente por {signer['name']} ({signer['email']}) en el portal Data Symphony, "
            f"{at.strftime('%Y-%m-%d %H:%M')} (hora de Puerto Rico)")
    return _write_cells(content, ws.title, edits)


def _write_cells(content, title, edits):
    """The workbook with the cells of sheet `title` set to text (see set_cell)."""
    source = zipfile.ZipFile(io.BytesIO(content))
    path = _sheet_path(source, title)
    xml = source.read(path).decode("utf-8")
    for ref, text in edits.items():
        xml = set_cell(xml, ref, text)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in source.infolist():
            z.writestr(item, xml.encode("utf-8") if item.filename == path else source.read(item.filename))
    return out.getvalue()


# ── database ─────────────────────────────────────────────────────────────────

_TABLES_READY = False


def _ensure_tables(cur, conn):
    global _TABLES_READY
    certs._ensure_tables(cur, conn)
    if _TABLES_READY:
        return
    have = certs._existing(cur, DB, list(_DDL))
    for name, ddl in _DDL.items():
        if name.upper() not in have:
            cur.execute(f"CREATE TABLE [{DB}].dbo.[{name}] ({ddl})")
    cols = certs._columns(cur, T_FORM)
    for col, sql_type in _ESIGN_COLUMNS:
        if col.upper() not in cols:
            cur.execute(f"ALTER TABLE [{DB}].dbo.[{T_FORM}] ADD [{col}] {sql_type} NULL")
    conn.commit()
    _TABLES_READY = True


_FORM_COLUMNS = ("[ID], [MOCK], [Source], [Agency], [BU], [Kind], [File_Name], [S3_Key], [Size_Bytes], "
                 "[Form_Agency], [Signer_Name], [Signer_Title], [Signed_Date], [Has_Signature_Image], "
                 "[Rows_Total], [Rows_Answered], [Uploaded_By], [Uploaded_DTTM], [Is_Current], [Revoked_By], "
                 "[Revoked_DTTM], [Revoke_Reason], [Sign_Method], [Signer_IP], [Signer_Agent], [Content_SHA256]")


def _form(r):
    return {"id": r[0], "mock": _s(r[1]), "source": _s(r[2]), "agency": _s(r[3]), "bu": _s(r[4]), "kind": _s(r[5]),
            "file_name": r[6], "s3_key": r[7], "size": r[8], "form_agency": r[9], "signer_name": r[10],
            "signer_title": r[11], "signed_date": r[12], "signature_image": bool(r[13]), "rows": r[14],
            "answered": r[15], "uploaded_by": r[16], "uploaded_at": r[17], "current": bool(r[18]),
            "revoked_by": r[19], "revoked_at": r[20], "revoke_reason": r[21], "electronic": r[22] == "ESIGN",
            "method": r[22] or "XLSX", "signer_ip": r[23], "signer_agent": r[24], "sha256": r[25]}


def _public(form):
    return {k: v for k, v in form.items() if k != "s3_key"}


def _forms(cur, mock, current=True):
    cur.execute(f"SELECT {_FORM_COLUMNS} FROM [{DB}].dbo.[{T_FORM}] WHERE [MOCK] = ?"
                + (" AND [Is_Current] = 1" if current else "") + " ORDER BY [ID]", (mock,))
    return [_form(r) for r in cur.fetchall()]


def _form_row(cur, raw_id):
    cur.execute(f"SELECT {_FORM_COLUMNS} FROM [{DB}].dbo.[{T_FORM}] WHERE [ID] = ?", (certs._id(raw_id),))
    r = cur.fetchone()
    if not r:
        raise ApiError("Form not found", 404)
    return _form(r)


def _rows_of(cur, form_id):
    cur.execute(
        f"SELECT [Row_Number], [Segment], [Data_Entity], [Folder], [Agency_Resource], [Comment], [Response_Code] "
        f"FROM [{DB}].dbo.[{T_ROW}] WHERE [Form_ID] = ? ORDER BY [Row_Number]", (form_id,))
    return [{"row": r[0], "segment": r[1], "entity": r[2], "folder": r[3], "resource": r[4], "comment": r[5],
             "code": r[6]} for r in cur.fetchall()]


# ── access ───────────────────────────────────────────────────────────────────

def _party_rows(rows, source, agency):
    wanted = authz.party_key(source, agency)
    return [r for r in rows if authz.party_key(r["source"], r["agency"]) == wanted]


def _kinds(party_rows):
    return sorted({certs.form_kind(party_rows, r) for r in party_rows})


def _kind(values, party_rows):
    kind = _s(values.get("kind")).upper()
    if kind not in KINDS:
        raise ApiError(f"kind must be one of {', '.join(KINDS)}")
    if kind not in _kinds(party_rows):
        raise ApiError(f"The {KINDS[kind]} form is not one this party certifies", 404)
    return kind


def _read_party(conn, cur, p, what):
    """(access, mock, the party's setup rows, source, agency, bu) for a read."""
    access = certs._reader(p.get("email"), what)
    mock = certs._mock(p.get("mock"))
    rows, source, agency, bu = certs._resolve_party(conn, cur, access, mock, certs._need(p, "source"),
                                                    _s(p.get("agency")), [])
    return access, mock, _party_rows(rows, source, agency), source, agency, bu


def _checked_form(cur, raw_id, email, what):
    access = certs._reader(email, what)
    form = _form_row(cur, raw_id)
    access.check(form["source"], form["agency"], form["bu"])
    return access, form


# ── templates (super users upload, everyone downloads) ───────────────────────

def esign_open(cur):
    """Whether agencies may sign in the portal (super users always may)."""
    return (app_settings.get_setting(cur, ESIGN_SETTING, "off")[0] or "") == "on"


def _templates(conn, cur, p, bucket):
    certs._reader(p.get("email"), "viewing the certification forms")
    mock = certs._mock(p.get("mock"))
    app_settings._ensure_tables(cur, conn)
    value, by, at = app_settings.get_setting(cur, ESIGN_SETTING, "off")
    s3 = boto3.client("s3")
    out = []
    for kind, label in KINDS.items():
        try:
            head = s3.head_object(Bucket=bucket, Key=template_key(mock, kind))
            out.append({"kind": kind, "label": label, "uploaded": True, "size": head["ContentLength"],
                        "last_modified": head["LastModified"].isoformat()})
        except Exception:  # noqa: BLE001 - no template for this kind yet
            out.append({"kind": kind, "label": label, "uploaded": False})
    return {"mock": mock, "templates": out, "esign_open": value == "on", "esign_changed_by": by, "esign_changed_at": at}


def _esign_setting(conn, cur, data, bucket):
    """A super user opens (or closes) electronic signature to the agencies."""
    actor = _s(data.get("actor"))
    authz.require(actor, what="changing electronic signature")
    app_settings._ensure_tables(cur, conn)
    app_settings.set_setting(cur, ESIGN_SETTING, "on" if data.get("open") is True else "off", actor)
    conn.commit()
    return {"esign_open": esign_open(cur)}


def _template_kind(values):
    kind = _s(values.get("kind")).upper()
    if kind not in KINDS:
        raise ApiError(f"kind must be one of {', '.join(KINDS)}")
    return kind


def _template_upload_url(data, bucket):
    authz.require(_s(data.get("actor")), what="replacing a certification form")
    mock, kind = certs._mock(data.get("mock")), _template_kind(data)
    name = certs._need(data, "file_name", "File name", 300)
    if not name.lower().endswith(".xlsx"):
        raise ApiError("The certification form must be an Excel workbook (.xlsx)")
    try:
        size = int(data.get("size") or 0)
    except (TypeError, ValueError):
        raise ApiError("size must be a number of bytes")
    if not 0 < size <= MAX_FORM_BYTES:
        raise ApiError(f"The form must be between 1 byte and {MAX_FORM_BYTES // (1024 * 1024)} MB")
    key = template_key(mock, kind)
    return {"key": key, "content_type": XLSX, "url": api_util.presign_put(bucket, key, XLSX)}


def _template_url(p, bucket):
    certs._reader(p.get("email"), "downloading a certification form")
    mock, kind = certs._mock(p.get("mock")), _template_kind(p)
    key = template_key(mock, kind)
    if api_util.object_size(bucket, key) is None:
        raise ApiError(f"The {KINDS[kind]} form has not been added for {mock} yet", 404)
    return {"url": api_util.presign_get(bucket, key, f"{mock}_{KINDS[kind]}_Certification_Template.xlsx")}


def _template_delete(data, bucket):
    authz.require(_s(data.get("actor")), what="removing a certification form")
    mock, kind = certs._mock(data.get("mock")), _template_kind(data)
    boto3.client("s3").delete_object(Bucket=bucket, Key=template_key(mock, kind))
    return {"kind": kind, "deleted": True}


# ── a party's forms ──────────────────────────────────────────────────────────

def _status(conn, cur, p, bucket):
    _, mock, mine, source, agency, bu = _read_party(conn, cur, p, "viewing the certification forms")
    party = authz.party_key(source, agency)
    current = {f["kind"]: f for f in _forms(cur, mock) if authz.party_key(f["source"], f["agency"]) == party}
    kinds = []
    for kind in _kinds(mine):
        covers = [r for r in mine if certs.form_kind(mine, r) == kind]
        form = current.get(kind)
        kinds.append({"kind": kind, "label": KINDS[kind],
                      "template": api_util.object_size(bucket, template_key(mock, kind)) is not None,
                      "records": [{f: r[f] for f in ("module", "file_type", "entity")} for r in covers],
                      "form": _public(form) if form else None,
                      "rows": _rows_of(cur, form["id"]) if form else []})
    return {"mock": mock, "source": source, "agency": agency, "bu": bu, "kinds": kinds, "esign_open": esign_open(cur)}


def _list(conn, cur, p, bucket):
    access = certs._reader(p.get("email"), "viewing the certification forms")
    mock = certs._mock(p.get("mock"))
    return {"forms": [_public(f) for f in _forms(cur, mock) if access.party(f["source"], f["agency"], f["bu"])]}


def _rows(conn, cur, p, bucket):
    _, form = _checked_form(cur, p.get("id"), p.get("email"), "viewing a certification form")
    return {"form": _public(form), "rows": _rows_of(cur, form["id"])}


def _file_url(conn, cur, p, bucket):
    _, form = _checked_form(cur, p.get("id"), p.get("email"), "downloading a certification form")
    view = bool(p.get("view")) and form["file_name"].lower().endswith(".pdf")
    return {"file_name": form["file_name"], "inline": view,
            "url": api_util.presign_get(bucket, form["s3_key"], form["file_name"], inline_type=PDF if view else None)}


def _download(conn, cur, data, bucket):
    """The cycle's blank form for one party, with its agency and source filled in."""
    access = certs._reader(data.get("actor"), "downloading a certification form")
    mock = certs._mock(data.get("mock"))
    rows, source, agency, bu = certs._resolve_party(conn, cur, access, mock, certs._need(data, "source"),
                                                    _s(data.get("agency")), [])
    kind = _kind(data, _party_rows(rows, source, agency))
    party = {"source": source, "agency": agency, "bu": bu}
    name = download_name(mock, kind, party)
    key = f"{FORMS_PREFIX}/{mock}/out/{uuid.uuid4().hex}/{name}"
    boto3.client("s3").put_object(Bucket=bucket, Key=key, Body=_party_form(bucket, mock, kind, party),
                                  ContentType=XLSX)
    return {"name": name, "url": api_util.presign_get(bucket, key, name)}


def _party_form(bucket, mock, kind, party):
    """The cycle's blank form of a kind with the party filled in."""
    try:
        content = boto3.client("s3").get_object(Bucket=bucket, Key=template_key(mock, kind))["Body"].read()
    except Exception:  # noqa: BLE001 - the template has not been uploaded
        raise ApiError(f"The {KINDS[kind]} form has not been added for {mock} yet. The validation team adds it "
                       "under User Guides.", 404)
    return fill_template(content, party)


def _upload_url(conn, cur, data, bucket):
    _, mock, rows, source, agency, bu = certs._writable_party(conn, cur, data, "uploading a certification form")
    kind = _kind(data, _party_rows(rows, source, agency))
    name = certs._need(data, "file_name", "File name", 300)
    if not name.lower().endswith((".xlsx", ".pdf")):
        raise ApiError("Upload the completed form as the Excel workbook (.xlsx) it was downloaded as, "
                       "or the signed form as a PDF")
    try:
        size = int(data.get("size") or 0)
    except (TypeError, ValueError):
        raise ApiError("size must be a number of bytes")
    if not 0 < size <= MAX_FORM_BYTES:
        raise ApiError(f"The form must be between 1 byte and {MAX_FORM_BYTES // (1024 * 1024)} MB")
    key = certs.attachment_key(mock, source, agency, bu, "FORM", kind, name)
    content_type = PDF if name.lower().endswith(".pdf") else XLSX
    return {"key": key, "content_type": content_type, "url": api_util.presign_put(bucket, key, content_type)}


def _submit(conn, cur, data, bucket):
    """Read an uploaded form, keep it, and certify what it covers."""
    access, mock, rows, source, agency, bu = certs._writable_party(conn, cur, data, "submitting a certification form")
    mine = _party_rows(rows, source, agency)
    kind = _kind(data, mine)
    key = certs._need(data, "key", limit=certs.MAX_KEY_LENGTH)
    name = certs._need(data, "file_name", "File name", 300)
    prefix = certs.cert_prefix(mock, source, agency, bu, "FORM", kind)
    if not key.startswith(prefix) or "/" in key[len(prefix):]:
        raise ApiError("The uploaded file does not belong to this form", 403)
    try:
        obj = boto3.client("s3").get_object(Bucket=bucket, Key=key)
    except Exception:  # noqa: BLE001
        raise ApiError(f"{name} was not uploaded", 404)
    if obj["ContentLength"] > MAX_FORM_BYTES:
        raise ApiError(f"The form is larger than {MAX_FORM_BYTES // (1024 * 1024)} MB")
    party = {"source": source, "agency": agency, "bu": bu}
    content = obj["Body"].read()
    if name.lower().endswith(".pdf"):
        return _submit_pdf(conn, cur, data, bucket, access, mock, party, mine, kind, key, name, content)
    form = read_form(content)
    errors, warnings = form_problems(form, party)
    if errors:
        raise ApiError(". ".join(errors) + ".")
    return _record(conn, cur, access, mock, party, mine, kind, key, name, obj["ContentLength"], form, warnings)


def _answers(data, table, kind):
    """{row number: (resource, comment text)} from the answers a page sent,
    checked against the form's rows and the comments each row offers."""
    answers = {}
    for a in data.get("answers") or []:
        try:
            row = int(a.get("row"))
        except (TypeError, ValueError, AttributeError):
            raise ApiError("Every answer needs the row of the form it answers")
        if row not in table:
            raise ApiError(f"Row {row} is not a data entity row of the {KINDS[kind]} form", 404)
        code = _s(a.get("code")).upper() or None
        if code and code not in certs.responses_for(table[row]["folder"] or "Validation"):
            raise ApiError(f"{table[row]['entity']}: that comment is not one of the options for this row")
        resource = _s(a.get("resource"))[:200]
        if code or resource:
            answers[row] = (resource, certs.RESPONSES[code] if code else "")
    return answers


def _submit_pdf(conn, cur, data, bucket, access, mock, party, mine, kind, key, name, content):
    """A signed PDF: kept as the form, with the answers the uploader recorded."""
    if not content.startswith(b"%PDF-"):
        raise ApiError(f"{name} is not a PDF file")
    signer = certs._need(data, "signer_name", "Name of the person who signed", 200)
    title = certs._need(data, "signer_title", "Title of the person who signed", 200)
    if data.get("confirm") is not True:
        raise ApiError("Confirm that the answers are those of the signed PDF")
    signed_on = certs._date(data.get("signed_date"), "Date signed") if data.get("signed_date") else None
    table = {x["row"]: x for x in read_form(_party_form(bucket, mock, kind, party))["rows"]}
    answers = _answers(data, table, kind)
    rows = []
    for number, x in table.items():
        resource, comment = answers.get(number, ("", ""))
        rows.append({**x, "resource": resource, "comment": comment, "code": match_comment(comment) if comment else None})
    form = {"agency": agency_label(party), "name": signer, "title": title,
            "date": signed_on.isoformat() if signed_on else date.today().isoformat(), "signature_image": False,
            "rows": rows, "pdf": True}
    errors, warnings = form_problems(form, party)
    if errors:
        raise ApiError(". ".join(errors) + ".")
    client = data.get("_client") or {}
    return _record(conn, cur, access, mock, party, mine, kind, key, name, len(content), form, warnings,
                   esign={"method": "PDF", "ip": client.get("ip"), "agent": client.get("agent"),
                          "sha256": hashlib.sha256(content).hexdigest(), "consent": PDF_STATEMENT})


_NOTE = {"ESIGN": "Signed electronically on the form", "PDF": "From the signed PDF"}


def _record(conn, cur, access, mock, party, mine, kind, key, name, size, form, warnings, esign=None):
    """Keep a read form as the party's current one and certify what it covers.
    esign: {ip, agent, sha256, consent} for a form signed in the portal."""
    source, agency, bu = party["source"], party["agency"], party["bu"]
    esign = esign or {}
    covers = [r for r in mine if certs.form_kind(mine, r) == kind]
    answers = assign(covers, form["rows"])

    cur.execute(f"SELECT [ID] FROM [{DB}].dbo.[{T_FORM}] WHERE [MOCK] = ? AND [Source] = ? AND [Agency] = ? "
                "AND [Kind] = ? AND [Is_Current] = 1", (mock, source, agency, kind))
    earlier = [r[0] for r in cur.fetchall()]
    answered = sum(1 for x in form["rows"] if x["code"])
    cur.execute(
        f"INSERT INTO [{DB}].dbo.[{T_FORM}] ([MOCK], [Source], [Agency], [BU], [Kind], [File_Name], [S3_Key], "
        "[Size_Bytes], [Form_Agency], [Signer_Name], [Signer_Title], [Signed_Date], [Has_Signature_Image], "
        "[Rows_Total], [Rows_Answered], [Uploaded_By], [Uploaded_DTTM], [Is_Current], [Sign_Method], [Signer_IP], "
        "[Signer_Agent], [Content_SHA256], [Consent_Text]) OUTPUT INSERTED.[ID] "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, GETDATE(), 1, ?, ?, ?, ?, ?)",
        (mock, source, agency, bu, kind, name, key, size, form["agency"][:200] or None,
         form["name"][:200], form["title"][:200] or None, form["date"][:50] or None, form["signature_image"],
         len(form["rows"]), answered, access.email, (esign.get("method") or "ESIGN") if esign else None,
         _s(esign.get("ip"))[:64] or None,
         _s(esign.get("agent"))[:400] or None, esign.get("sha256"), esign.get("consent")))
    form_id = cur.fetchone()[0]
    cur.fast_executemany = True
    cur.executemany(
        f"INSERT INTO [{DB}].dbo.[{T_ROW}] ([Form_ID], [Row_Number], [Segment], [Data_Entity], [Folder], "
        "[Agency_Resource], [Comment], [Response_Code]) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        [(form_id, x["row"], x["segment"][:100], x["entity"][:200], x["folder"][:100], x["resource"][:200],
          x["comment"][:1000], x["code"]) for x in form["rows"]])
    cur.fast_executemany = False
    if earlier:
        marks = ", ".join("?" for _ in earlier)
        cur.execute(f"UPDATE [{DB}].dbo.[{T_FORM}] SET [Is_Current] = 0 WHERE [ID] IN ({marks})", tuple(earlier))
        # What the earlier form certified gives way to this one.
        cur.execute(f"UPDATE [{DB}].dbo.[{certs.T_FILE}] SET [Is_Current] = 0 WHERE [Form_ID] IN ({marks}) "
                    "AND [Is_Current] = 1", tuple(earlier))

    # Rows answered as incorrect become issues. An issue the earlier form
    # already raised for the same row keeps its documents.
    kept = {}
    if earlier:
        cur.execute(f"SELECT [ID], [Description] FROM [{DB}].dbo.[{certs.T_ISSUE}] WHERE [Deleted] = 0 AND "
                    f"[Form_ID] IN ({', '.join('?' for _ in earlier)})", tuple(earlier))
        kept = {r[1]: r[0] for r in cur.fetchall()}
    by_key = {certs._row_key(r): r for r in covers}
    raised, issues = set(), []
    for record_key, (code, chosen) in answers.items():
        rec = by_key[record_key]
        resources = "; ".join(dict.fromkeys(x["resource"] for x in chosen if x["resource"]))[:200] or form["name"][:200]
        cur.execute(
            f"UPDATE [{DB}].dbo.[{certs.T_FILE}] SET [Is_Current] = 0 WHERE [MOCK] = ? AND [Source] = ? "
            "AND [Agency] = ? AND [Entity] = ? AND [File_Type] = ? AND ([Module] = ? OR [Module] IS NULL) "
            "AND [Is_Current] = 1", (mock, source, agency, rec["entity"], rec["file_type"], rec["module"]))
        cur.execute(
            f"INSERT INTO [{DB}].dbo.[{certs.T_FILE}] ([MOCK], [Source], [Agency], [BU], [Module], [Entity], "
            "[File_Type], [Response_Code], [Resource_Name], [Notes], [Certified_By], [Certified_DTTM], "
            "[Is_Current], [Form_ID]) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, GETDATE(), 1, ?)",
            (mock, source, agency, bu, rec["module"], rec["entity"], rec["file_type"], code, resources,
             f"{_NOTE.get(esign.get('method') or 'ESIGN') if esign else 'From the signed form'} {name}", access.email,
             form_id))
        for x in chosen:
            if x["code"] != "ISSUES" or id(x) in raised:
                continue
            raised.add(id(x))
            description = f"{x['entity']} ({x['folder'] or 'Validation'}) - {x['segment']}".strip(" -")
            if description in kept:
                cur.execute(f"UPDATE [{DB}].dbo.[{certs.T_ISSUE}] SET [Form_ID] = ? WHERE [ID] = ?",
                            (form_id, kept.pop(description)))
                continue
            cur.execute(
                f"INSERT INTO [{DB}].dbo.[{certs.T_ISSUE}] ([MOCK], [Source], [Agency], [Module], [File_Type], "
                "[Entity], [Description], [Reported_By], [Reported_DTTM], [Deleted], [Form_ID]) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, GETDATE(), 0, ?)",
                (mock, source, agency, rec["module"], rec["file_type"], rec["entity"], description, access.email,
                 form_id))
            issues.append(description)
    if kept:
        # Rows no longer answered as incorrect.
        cur.execute(f"UPDATE [{DB}].dbo.[{certs.T_ISSUE}] SET [Deleted] = 1 WHERE [ID] IN "
                    f"({', '.join('?' for _ in kept)})", tuple(kept.values()))
    conn.commit()
    missing = [{f: r[f] for f in ("module", "file_type", "entity")} for r in covers
               if certs._row_key(r) not in answers]
    if missing:
        warnings.append(f"{len(missing)} record(s) this form covers have no answered row and stay pending")
    return {"form": _public(_form_row(cur, form_id)), "rows": _rows_of(cur, form_id),
            "certified": len(answers), "not_certified": missing, "new_issues": issues, "warnings": warnings}


def _esign_form(conn, cur, p, bucket):
    """The rows of a party's form to answer in the portal, with the comments
    each row offers and the answers of the party's current form."""
    access, mock, mine, source, agency, bu = _read_party(conn, cur, p, "signing a certification form")
    kind = _kind(p, mine)
    party = {"source": source, "agency": agency, "bu": bu}
    rows = read_form(_party_form(bucket, mock, kind, party))["rows"]
    current = next((f for f in _forms(cur, mock) if f["kind"] == kind
                    and authz.party_key(f["source"], f["agency"]) == authz.party_key(source, agency)), None)
    earlier = {(_plain(x["segment"]), _entity_key(x["entity"]), _plain(x["folder"])): x
               for x in (_rows_of(cur, current["id"]) if current else [])}
    out = []
    for x in rows:
        before = earlier.get((_plain(x["segment"]), _entity_key(x["entity"]), _plain(x["folder"]))) or {}
        out.append({"row": x["row"], "segment": x["segment"], "entity": x["entity"], "folder": x["folder"],
                    "options": list(certs.responses_for(x["folder"] or "Validation")),
                    "resource": _s(before.get("resource")), "code": before.get("code")})
    perms = authz.get_permissions(access.email)
    return {"kind": kind, "label": KINDS[kind], "agency_label": agency_label(party), "rows": out,
            "responses": certs.RESPONSES, "consent": ESIGN_CONSENT, "esign_open": esign_open(cur),
            "signer": {"email": access.email, "name": _s(perms.get("name"))}}


def _esign(conn, cur, data, bucket):
    """Answer and sign a form in the portal. The portal writes the answers and
    the signature into the cycle's form, keeps that file and records it like
    an uploaded one."""
    access, mock, rows, source, agency, bu = certs._writable_party(conn, cur, data, "signing a certification form")
    if access.role != authz.SUPER_USER and not esign_open(cur):
        raise ApiError("Electronic signature is not available yet. Download the form, sign it and upload it.", 403)
    mine = _party_rows(rows, source, agency)
    kind = _kind(data, mine)
    name = certs._need(data, "signer_name", "Name", 200)
    title = certs._need(data, "signer_title", "Title", 200)
    if data.get("consent") is not True:
        raise ApiError("Confirm that you sign this form electronically")
    party = {"source": source, "agency": agency, "bu": bu}
    blank = _party_form(bucket, mock, kind, party)
    table = {x["row"]: x for x in read_form(blank)["rows"]}
    answers = _answers(data, table, kind)
    at = datetime.now(_PR_TIME)
    content = sign_form(blank, answers, {"name": name, "title": title, "email": access.email, "at": at})
    form = read_form(content)
    form["electronic"] = True
    errors, warnings = form_problems(form, party)
    if errors:
        raise ApiError(". ".join(errors) + ".")
    file_name = download_name(mock, kind, party)[:-len(".xlsx")] + "_eSigned.xlsx"
    key = certs.attachment_key(mock, source, agency, bu, "FORM", kind, file_name)
    boto3.client("s3").put_object(Bucket=bucket, Key=key, Body=content, ContentType=XLSX)
    client = data.get("_client") or {}
    return _record(conn, cur, access, mock, party, mine, kind, key, file_name, len(content), form, warnings,
                   esign={"ip": client.get("ip"), "agent": client.get("agent"),
                          "sha256": hashlib.sha256(content).hexdigest(), "consent": ESIGN_CONSENT})


def _revoke(conn, cur, data, bucket):
    """A super user sends a form back: it and what it certified stop counting."""
    actor = _s(data.get("actor"))
    authz.require(actor, what="rejecting a certification form")
    form = _form_row(cur, data.get("id"))
    reason = certs._need(data, "reason", "Reason")
    if not form["current"]:
        raise ApiError("This form has already been replaced or rejected", 409)
    certs._check_open(cur, form["mock"], form["source"], form["agency"])
    cur.execute(f"UPDATE [{DB}].dbo.[{T_FORM}] SET [Is_Current] = 0, [Revoked_By] = ?, [Revoked_DTTM] = GETDATE(), "
                "[Revoke_Reason] = ? WHERE [ID] = ?", (actor, reason, form["id"]))
    cur.execute(f"UPDATE [{DB}].dbo.[{certs.T_FILE}] SET [Is_Current] = 0, [Revoked_By] = ?, "
                "[Revoked_DTTM] = GETDATE(), [Revoke_Reason] = ? WHERE [Form_ID] = ? AND [Is_Current] = 1",
                (actor, reason, form["id"]))
    revoked = cur.rowcount
    conn.commit()
    return {"id": form["id"], "revoked": True, "records": revoked}


_S3_ONLY = {"certform_template_upload_url": _template_upload_url,
            "certform_template_url": _template_url, "certform_template_delete": _template_delete}
_WITH_DATABASE = {"certform_templates": _templates, "certform_status": _status, "certform_list": _list,
                  "certform_rows": _rows, "certform_file_url": _file_url, "certform_download": _download,
                  "certform_upload_url": _upload_url, "certform_submit": _submit, "certform_revoke": _revoke,
                  "certform_esign_form": _esign_form, "certform_esign": _esign,
                  "certform_esign_setting": _esign_setting}


def handle(action, event, bucket, headers, conn_str):
    def run():
        values = api_util.params(event) if action in _READS else api_util.body(event)
        if action in ("certform_esign", "certform_submit"):
            # Where the signature came from, for the record of an electronic signature.
            http = (event.get("requestContext") or {}).get("http") or {}
            values["_client"] = {"ip": http.get("sourceIp"), "agent": http.get("userAgent")}
        if action in _S3_ONLY:
            return api_util.ok(headers, _S3_ONLY[action](values, bucket))
        with pyodbc.connect(conn_str, autocommit=False) as conn:
            cur = conn.cursor()
            _ensure_tables(cur, conn)
            return api_util.ok(headers, _WITH_DATABASE[action](conn, cur, values, bucket))

    return api_util.guarded(run, headers)
