"""
validation_report.py — the Excel workbooks validation hands back to the
client, in the layouts the validation team's programs produce.

Run workbook (FIN / SCM programs):
  Data            one row per failing record — Error Message, Entity, File,
                  Source, Last Load Dttm, Validation Type, Column 1 .. Column 21
  Summary         count of rows by Validation Type > Entity > Error Message,
                  with a total per Validation Type and a Grand Total
  Error Messages  the catalog entry of every validation code present

HCM workbook (one program run, or every reportable HCM program of a source /
business unit — the per-agency workbook):
  Summary               as above
  Observaciones         left for the reviewer's notes
  File Validation Error one row per failing record — the ten named columns
                        plus Column 1 .. Column 30
  Error Messages        as above

Files are written to S3 under REPORT_PREFIX/<mock>/<program>/<source>/ with
the team's file naming, e.g. 122651_HUM_ASSETS_MOCK14_FileValidation_2026.09.19-14.05.31.xlsx,
and the per-agency workbooks under REPORT_PREFIX/<mock>/Agency/<source>/.

The HCM workbooks carry person numbers, so downloads go through the
val_report_url action, which checks the caller's role before handing out a
time-limited link.
"""
import io
import re
from collections import Counter
from datetime import datetime

import boto3
import pyodbc
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import api_util
import authz
import validation_seed
from api_util import ApiError

ACTIONS = {"val_report_url", "val_agency_report"}

# Where validation data lives (the same database the runner works against).
DB = validation_seed.TARGET_DB

# Under the private root: the workbooks carry person numbers, so the browser's
# storage rights must not reach them. Downloads go through val_report_url.
REPORT_PREFIX = f"{api_util.PRIVATE_ROOT}/Reports"
AGENCY_DIR = "Agency"
N_EXTRA = 21
HEADER = (["Error Message", "Entity", "File", "Source", "Last Load Dttm", "Validation Type"]
          + [f"Column {i}" for i in range(1, N_EXTRA + 1)])

# HCM layout: the sheet header and the LOG_DATA_CLEANSE_DETAIL column behind
# each position, in the same order.
N_HCM_EXTRA = 30
HCM_HEADER = (["Error Message", "Entity", "File", "Last Load Dttm", "Validation Type", "Source",
               "Validation Code", "Mock", "Person Number", "Business Unit"]
              + [f"Column {i}" for i in range(1, N_HCM_EXTRA + 1)])
HCM_DETAIL_COLS = (["ERROR_MSG", "Entity", "File", "File_PROCESSED_DTTM", "VALIDATION_TYPE", "Source",
                    "Validation_Code", "MOCK", "PersonNumber", "BU"]
                   + [f"Col{i}" for i in range(1, N_HCM_EXTRA + 1)])
_HCM_MOCK_IDX = HCM_HEADER.index("Mock")
LEGEND_HEADER = ["Validation Code", "Error Message", "Error Message Description", "Entity", "Validation Type"]
SUMMARY_HEADER = ["Row Labels", "Count of Error Message"]
OBSERVACIONES_TITLE = "Observaciones DCL y Validaciones"

# The team's programs name the workbook by entity, not by catalog program.
_LABEL = {"Asset": "ASSETS", "Inventory": "ITEM"}
# Catalog flag that marks a rule as reported to each audience.
_AUDIENCE_FLAG = {"agency": "AgencyReports", "source": "SourceReports"}
_URL_PREFIXES = (f"{REPORT_PREFIX}/", f"{api_util.PRIVATE_ROOT}/DataCleanseLog/")

_HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_BOLD = Font(bold=True)


def _fixed(cols, row, name_variants):
    """Value of the first column whose (normalised) name is in name_variants."""
    for i, c in enumerate(cols):
        if c.upper().replace(" ", "_") in name_variants:
            return row[i]
    return ""


def records_from_view_rows(cols, rows):
    """Map one view's rows onto the team's 27-column layout: the six named
    columns by name, every other column in view order into Column 1..21."""
    fixed_idx = set()
    order = [("ERROR_MSG",), ("ENTITY",), ("FILE",), ("SOURCE",),
             ("PROCESSED_DTTM", "LAST_LOAD_DTTM", "FILE_PROCESSED_DTTM", "LAST_LOAD_DTTM"),
             ("VALIDATION_TYPE",)]
    picks = []
    for variants in order:
        idx = next((i for i, c in enumerate(cols) if c.upper().replace(" ", "_") in variants), None)
        picks.append(idx)
        if idx is not None:
            fixed_idx.add(idx)
    extra_idx = [i for i in range(len(cols)) if i not in fixed_idx][:N_EXTRA]
    out = []
    for r in rows:
        rec = [(r[i] if i is not None else "") for i in picks]
        rec += [r[i] for i in extra_idx]
        rec += [""] * (N_EXTRA - len(extra_idx))
        out.append(rec)
    return out


def records_from_detail_rows(rows):
    """LOG_DATA_CLEANSE_DETAIL rows (ERROR_MSG, Entity, File, File_PROCESSED_DTTM,
    Source, VALIDATION_TYPE, Col1..Col30) onto the same layout."""
    out = []
    for r in rows:
        rec = [r[0], r[1], r[2], r[4], r[3], r[5]]
        extras = [c for c in r[6:6 + 30] if c not in (None, "")][:N_EXTRA]
        rec += extras + [""] * (N_EXTRA - len(extras))
        out.append(rec)
    return out


# ── sheets ───────────────────────────────────────────────────────────────────

def _styled(ws, value, bold=False, indent=0, head=False):
    safe = api_util.xlsx_value(ws, value)
    cell = safe if isinstance(safe, Cell) else WriteOnlyCell(ws, value=safe)
    if head:
        cell.fill = _HEAD_FILL
        cell.font = _HEAD_FONT
    elif bold:
        cell.font = _BOLD
    if indent:
        cell.alignment = Alignment(indent=indent)
    return cell


def _group_label(value):
    return str(value if value is not None else "").strip() or "(blank)"


def _write_summary(ws, triples):
    """Summary sheet from (validation type, entity, error message) per record:
    rows grouped Validation Type > Entity > Error Message with their counts."""
    ws.column_dimensions["A"].width = 100
    ws.column_dimensions["B"].width = 24
    ws.freeze_panes = "A2"
    ws.append([_styled(ws, h, head=True) for h in SUMMARY_HEADER])
    tree = {}
    for vtype, entity, msg in triples:
        by_entity = tree.setdefault(_group_label(vtype), {})
        by_entity.setdefault(_group_label(entity), Counter())[_group_label(msg)] += 1
    grand = 0
    for vtype in sorted(tree):
        subtotal = sum(sum(msgs.values()) for msgs in tree[vtype].values())
        grand += subtotal
        ws.append([_styled(ws, vtype, bold=True), _styled(ws, subtotal, bold=True)])
        for entity in sorted(tree[vtype]):
            msgs = tree[vtype][entity]
            ws.append([_styled(ws, entity, bold=True, indent=1), _styled(ws, sum(msgs.values()), bold=True)])
            for msg in sorted(msgs):
                ws.append([_styled(ws, msg, indent=2), msgs[msg]])
    ws.append([_styled(ws, "Grand Total", bold=True), _styled(ws, grand, bold=True)])


def _write_legend(ws, legend):
    """Error Messages sheet: one row per validation code."""
    for i, w in enumerate([18, 70, 90, 24, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.append([_styled(ws, h, head=True) for h in LEGEND_HEADER])
    for row in legend:
        ws.append([api_util.xlsx_value(ws, _cell(v)) for v in list(row)[:len(LEGEND_HEADER)]])


def build_workbook(records, legend=None):
    """Workbook bytes: Data sheet + Summary sheet (+ Error Messages when
    legend rows are supplied)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(HEADER)
    for c in range(1, len(HEADER) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEAD_FILL
        cell.font = _HEAD_FONT
    for rec in records:
        ws.append([api_util.xlsx_value(ws, _cell(v)) for v in rec])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}{max(2, len(records) + 1)}"
    widths = [60, 18, 40, 14, 20, 14] + [22] * N_EXTRA
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    _write_summary(wb.create_sheet("Summary"), ((rec[5], rec[1], rec[0]) for rec in records))
    if legend:
        _write_legend(wb.create_sheet("Error Messages"), legend)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_hcm_workbook(records, legend, mock):
    """HCM workbook bytes: Summary, Observaciones, File Validation Error,
    Error Messages. `records` are rows in HCM_HEADER / HCM_DETAIL_COLS order;
    `mock` fills the Mock column of a record that has none. Written in
    streaming mode — these workbooks can run to hundreds of thousands of rows."""
    wb = Workbook(write_only=True)
    _write_summary(wb.create_sheet("Summary"), ((rec[4], rec[1], rec[0]) for rec in records))

    obs = wb.create_sheet("Observaciones")
    obs.column_dimensions["A"].width = 60
    obs.append([_styled(obs, OBSERVACIONES_TITLE, bold=True)])

    ws = wb.create_sheet("File Validation Error")
    widths = [60, 18, 40, 20, 16, 14, 16, 14, 16, 14] + [22] * N_HCM_EXTRA
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HCM_HEADER))}{max(2, len(records) + 1)}"
    ws.append([_styled(ws, h, head=True) for h in HCM_HEADER])
    for rec in records:
        row = [_hcm_cell(ws, v) for v in rec]
        if row[_HCM_MOCK_IDX] is None:
            row[_HCM_MOCK_IDX] = mock
        ws.append(row)

    _write_legend(wb.create_sheet("Error Messages"), legend or [])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    return v


def _hcm_cell(ws, v):
    """Empty values stay out of the streamed sheet altogether; text is made
    safe for the sheet (no control characters, never a live formula)."""
    v = _cell(v)
    return None if v == "" else api_util.xlsx_value(ws, v)


# ── S3 ───────────────────────────────────────────────────────────────────────

def _stamp(when=None):
    return (when or datetime.now()).strftime("%Y.%m.%d-%H.%M.%S")


def _program_dir(program):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", program)


def report_key(mock, program, source, when=None):
    label = _LABEL.get(program, re.sub(r"[^A-Za-z0-9]+", "_", program).strip("_").upper())
    name = f"{source}_{label}_{mock}_FileValidation_{_stamp(when)}.xlsx"
    return f"{REPORT_PREFIX}/{mock}/{_program_dir(program)}/{source}/{name}"


def hcm_report_key(mock, program, source, when=None):
    """Key of one HCM program run's workbook."""
    name = f"HCM_FileValidation_{source}_{_program_dir(program).strip('_')}_{mock}_{_stamp(when)}.xlsx"
    return f"{REPORT_PREFIX}/{mock}/{_program_dir(program)}/{source}/{name}"


def write_report(bucket, key, content):
    boto3.client("s3").put_object(
        Bucket=bucket, Key=key, Body=content,
        ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    return key


def list_reports(bucket, mock, program=None, source=None, include_agency=False):
    """Reports in S3 for a mock (optionally one program / source), newest first.
    With include_agency (an HCM program) the listing also carries the
    per-agency workbooks (program "Agency")."""
    base = f"{REPORT_PREFIX}/{mock}/"
    prefixes = [base]
    if program:
        tail = f"{source}/" if source else ""
        prefixes = [f"{base}{_program_dir(program)}/{tail}"]
        if include_agency:
            prefixes.append(f"{base}{AGENCY_DIR}/{tail}")
    s3 = boto3.client("s3")
    out = []
    for prefix in prefixes:
        for page in s3.get_paginator("list_objects_v2").paginate(Bucket=bucket, Prefix=prefix):
            for o in page.get("Contents", []):
                if not o["Key"].lower().endswith(".xlsx"):
                    continue
                parts = o["Key"].split("/")
                out.append({"key": o["Key"], "name": parts[-1], "program": parts[3] if len(parts) > 4 else "",
                            "source": parts[4] if len(parts) > 5 else "", "size": o["Size"],
                            "last_modified": o["LastModified"].isoformat()})
    out.sort(key=lambda r: r["last_modified"], reverse=True)
    return out


# ── catalog + per-agency workbook ────────────────────────────────────────────

def legend_rows(cur, codes, where, params=()):
    """Catalog rows (LEGEND_HEADER order) for the validation codes present.
    `where` is fixed SQL text over SETUP_ERROR_MESSAGES_SOURCE supplied by the
    caller, never user input."""
    wanted = {str(c or "").strip().upper() for c in codes}
    cur.execute(
        f"SELECT VALIDATION_CODE, ERROR_MESSAGE, ERROR_MESSAGE_LONG_DESCRIPTION, ENTITY, VALIDATION_TYPE "
        f"FROM [{DB}].dbo.SETUP_ERROR_MESSAGES_SOURCE WHERE {where} ORDER BY VALIDATION_CODE",
        list(params),
    )
    out, seen = [], set()
    for r in cur.fetchall():
        code = str(r[0] or "").strip().upper()
        if code in wanted and code not in seen:
            seen.add(code)
            out.append(list(r))
    return out


def _exists(cur, name):
    cur.execute(f"SELECT COUNT(*) FROM [{DB}].sys.objects WHERE name = ?", (name,))
    return cur.fetchone()[0] > 0


def _ensure_objects(conn_str, names, warnings):
    """Against a test database, clone whichever of `names` it does not have yet."""
    if not validation_seed.is_test_target():
        return
    with pyodbc.connect(conn_str, autocommit=False) as conn:
        cur = conn.cursor()
        missing = [n for n in names if not _exists(cur, n)]
        if not missing:
            return
        try:
            seeder = validation_seed.Seeder(conn)
            for name in missing:
                seeder.ensure(name)
            warnings.extend(f"Not created in {DB}: {f}" for f in seeder.report["failed"])
        except Exception as e:  # noqa: BLE001 - reported to the caller as a warning
            warnings.append(f"Could not prepare {', '.join(missing)} in {DB}: {str(e)[:200]}")


def _agency_name(cur, mock, source, bu):
    """The agency a source (+ business unit) belongs to, when the mock has a
    file-location setup table and it names exactly one."""
    table = f"SETUP_DATA_CLEANSE_FILE_LOCATION_{mock}"
    if not _exists(cur, table):
        return ""
    sql = (f"SELECT DISTINCT LTRIM(RTRIM([Agency])) FROM [{DB}].dbo.[{table}] "
           f"WHERE [Source] = ? AND ISNULL(LTRIM(RTRIM([Agency])), '') <> ''")
    params = [source]
    if bu:
        sql += " AND LEFT([BU], 3) = LEFT(?, 3)"
        params.append(bu)
    cur.execute(sql, params)
    names = [r[0] for r in cur.fetchall()]
    return names[0] if len(names) == 1 else ""


def agency_report(conn_str, bucket, mock, source, bu=None, audience="agency"):
    """Build and store the per-agency HCM workbook: every stored failing row of
    a source (+ business unit) for the HCM rules flagged as reported to the
    audience, whichever program produced them. Returns {key, name, rows, warnings}."""
    mock = api_util.mock(mock)
    source = api_util.ident(str(source or "").upper(), "source")
    bu = str(bu or "").strip()
    if bu and not re.fullmatch(r"[A-Za-z0-9]{1,20}", bu):
        raise ApiError(f"Invalid business unit: {bu!r}")
    flag = _AUDIENCE_FLAG.get(str(audience or "agency").strip().lower())
    if not flag:
        raise ApiError("audience must be 'agency' or 'source'")

    warnings = []
    _ensure_objects(conn_str, ["LOG_DATA_CLEANSE_DETAIL", "SETUP_ERROR_MESSAGES_SOURCE",
                               f"SETUP_DATA_CLEANSE_FILE_LOCATION_{mock}"], warnings)
    reported = f"Pillar = 'HCM' AND [{flag}] = 'Y'"
    sql = (f"SELECT {', '.join(f'd.[{c}]' for c in HCM_DETAIL_COLS)} "
           f"FROM [{DB}].dbo.LOG_DATA_CLEANSE_DETAIL d "
           f"WHERE d.MOCK = ? AND d.[Source] = ? AND d.Validation_Code IN ("
           f"SELECT VALIDATION_CODE FROM [{DB}].dbo.SETUP_ERROR_MESSAGES_SOURCE WHERE {reported})")
    params = [mock, source]
    if bu:
        sql += " AND LEFT(d.BU, 3) = LEFT(?, 3)"
        params.append(bu)
    sql += " ORDER BY d.Validation_Code, d.ERROR_MSG"
    with pyodbc.connect(conn_str) as conn:
        cur = conn.cursor()
        cur.execute(sql, params)
        records = cur.fetchall()
        if not records:
            raise ApiError(f"No reportable HCM validation rows stored for {source}"
                           f"{' / ' + bu if bu else ''} on {mock}", 404)
        code_idx = HCM_DETAIL_COLS.index("Validation_Code")
        legend = legend_rows(cur, {r[code_idx] for r in records}, reported)
        try:
            agency = _agency_name(cur, mock, source, bu)
        except Exception as e:  # noqa: BLE001 - the name is only part of the file name
            agency = ""
            warnings.append(f"Agency name not resolved: {str(e)[:200]}")

    who = "-".join(api_util.safe_segment(p) for p in (source, bu, agency) if p)
    name = f"HCM_FileValidation_{who}_{mock}_{_stamp()}.xlsx"
    key = f"{REPORT_PREFIX}/{mock}/{AGENCY_DIR}/{source}/{name}"
    write_report(bucket, key, build_hcm_workbook(records, legend, mock))
    return {"key": key, "name": name, "rows": len(records), "warnings": warnings}


# ── actions ──────────────────────────────────────────────────────────────────

def _key_codes(parts):
    """Source / business-unit codes a report key is filed under."""
    if parts[1] != "Reports":
        # The cleanse-log folders belong to their own module: any folder of
        # the key may be the source.
        return parts[2:-1]
    if len(parts) < 6:
        return []
    source, name = parts[4], parts[-1]
    codes = [source]
    if parts[3] == AGENCY_DIR:
        m = re.match(rf"^HCM_FileValidation_{re.escape(source)}-([A-Za-z0-9]+)[-_]", name)
        if m:
            codes.append(m.group(1))
    return codes


def _check_download(email, key):
    """403 unless the caller's role covers this key."""
    if not key.startswith(_URL_PREFIXES) or ".." in key or "\\" in key or key.endswith("/"):
        raise ApiError("Not a validation report key")
    role = authz.role_of(email)
    if role == authz.SUPER_USER:
        return
    parts = key.split("/")
    if role == authz.CERT_REVIEWER and "Certification" in parts[2:-1]:
        return
    if role == authz.AGENCY_USER and authz.can_act_on(email, *_key_codes(parts)):
        return
    raise ApiError(f"{email or 'Anonymous user'} is not allowed to download this report", 403)


def handle(action, event, bucket, headers, conn_str):
    def run():
        if action == "val_report_url":
            p = api_util.params(event)
            key = (p.get("key") or "").strip()
            _check_download((p.get("email") or "").strip(), key)
            if api_util.object_size(bucket, key) is None:
                raise ApiError("Report not found", 404)
            return api_util.ok(headers, {"url": api_util.presign_get(bucket, key, key.rsplit("/", 1)[-1])})

        body = api_util.body(event)
        authz.require((body.get("actor") or "").strip(), what="generating agency workbooks")
        res = agency_report(conn_str, bucket, body.get("mock"), body.get("source"),
                            bu=body.get("bu"), audience=body.get("audience") or "agency")
        res["url"] = api_util.presign_get(bucket, res["key"], res["name"])
        return api_util.ok(headers, res)

    return api_util.guarded(run, headers)
